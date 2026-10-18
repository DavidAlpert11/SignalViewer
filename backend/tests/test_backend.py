"""Backend tests: python -m pytest backend/tests"""
import base64
import json
import os
import sys

import numpy as np
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.store import store  # noqa: E402
from svcore.decimate import _minmax_indices, decimate_for_view  # noqa: E402
from svcore.mat_loader import load_mat  # noqa: E402

client = TestClient(app)


@pytest.fixture
def csv_file(tmp_path):
    t = np.arange(10000) / 1000
    p = tmp_path / "sig.csv"
    with open(p, "w", encoding="utf-8-sig") as f:   # BOM like Excel
        f.write("Time,Sine,Text\n")
        for i, ti in enumerate(t):
            f.write(f"{ti:.4f},{np.sin(ti):.6f},x{i}\n")
    yield str(p)
    for rid in list(store.order):
        store.remove(rid)


def test_open_and_data(csv_file):
    r = client.post("/api/files/open", json={"paths": [csv_file]}).json()
    assert not r["errors"]
    run = r["opened"][0]
    assert [s["name"] for s in run["signals"]] == ["Sine"]      # text column skipped
    sid = run["signals"][0]["id"]
    d = client.post("/api/data", json={"signals": [sid], "points": 1000}).json()["series"][0]
    assert d["decimated"] and d["total"] == 10000
    x = np.frombuffer(base64.b64decode(d["x"]), np.float64)
    assert np.all(np.diff(x) >= 0)


def test_missing_file_is_reported_not_raised():
    r = client.post("/api/files/open", json={"paths": ["/definitely/missing.csv"]}).json()
    assert r["errors"] and not r["opened"]


def test_cursor_and_stats_use_full_resolution(csv_file):
    run = client.post("/api/files/open", json={"paths": [csv_file]}).json()["opened"][0]
    sid = run["signals"][0]["id"]
    v = client.post("/api/cursor", json={"signals": [sid], "t": 1.5708}).json()["values"][sid]
    assert abs(v - 1.0) < 1e-3
    st = client.post("/api/stats", json={"signals": [sid], "t0": 0, "t1": 3.2}).json()["stats"][sid]
    assert abs(st["max"] - 1.0) < 1e-4


def test_ids_are_stable_across_remove_and_reload(csv_file, tmp_path):
    other = tmp_path / "b.csv"
    other.write_text("Time,A\n0,1\n1,2\n")
    a = client.post("/api/files/open", json={"paths": [csv_file]}).json()["opened"][0]
    b = client.post("/api/files/open", json={"paths": [str(other)]}).json()["opened"][0]
    client.delete(f"/api/runs/{a['id']}")
    assert client.post(f"/api/runs/{b['id']}/reload").json()["id"] == b["id"]
    assert [r["id"] for r in client.get("/api/runs").json()] == [b["id"]]


def test_restore_twice_does_not_duplicate(csv_file):
    run = client.post("/api/files/open", json={"paths": [csv_file]}).json()["opened"][0]
    payload = [{"id": run["id"], "path": run["path"], "name": run["name"], "importSettings": run["importSettings"]}]
    client.post("/api/session/restore", json={"runs": payload})
    client.post("/api/session/restore", json={"runs": payload})
    assert len(client.get("/api/runs").json()) == 1


def test_stream_append(csv_file):
    run = client.post("/api/files/open", json={"paths": [csv_file]}).json()["opened"][0]
    with open(csv_file, "a") as f:
        f.write("10.0,0.5,y\n10.001,0.6,z\n")
    r = client.post("/api/stream/poll").json()
    assert run["id"] in r["changed"]
    assert client.get("/api/runs").json()[0]["samples"] == 10002


def test_decimation_keeps_spike_and_has_no_coarse_tail():
    n = 6000
    y = np.random.default_rng(0).standard_normal(n)
    y[4321] = 50
    idx = _minmax_indices(y, 1600)
    assert 4321 in idx
    assert np.diff(idx).max() <= 2 * int(np.ceil(n / 800))     # remainder bin bug
    x = np.arange(n, dtype=float)
    xd, yd, dec = decimate_for_view(x, y, [1000, 1100], 4000)
    assert ((xd >= 1000) & (xd <= 1100)).sum() == 101           # raw samples when zoomed in


def test_mat_v5_and_v73(tmp_path):
    scipy_io = pytest.importorskip("scipy.io")
    h5py = pytest.importorskip("h5py")
    t = np.arange(0, 1, 0.01)
    p5 = tmp_path / "a.mat"
    scipy_io.savemat(p5, {"t": t, "v": np.sin(t), "s": {"w": np.cos(t)}, "m": np.column_stack([t, t])})
    run, skipped = load_mat(str(p5), "a.mat")
    assert set(run.signals) == {"v", "s.w", "m[1]", "m[2]"} and run.sample_count == 100
    p73 = tmp_path / "b.mat"
    with h5py.File(p73, "w") as f:
        f["time"] = t.reshape(1, -1)
        f["volt"] = np.sin(t).reshape(1, -1)
    run, _ = load_mat(str(p73), "b.mat")
    assert list(run.signals) == ["volt"]


# ------------------------------------------------------------------ phase 3

def _open(tmp_path, name, header, cols):
    p = tmp_path / name
    np.savetxt(p, np.column_stack(cols), delimiter=",", header=",".join(header), comments="")
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    return {s["name"]: s for s in run["signals"]}, run


def test_derived_integral_nonuniform_and_formula(tmp_path):
    t = np.cumsum(np.random.default_rng(1).uniform(0.005, 0.02, 2000))
    sigs, run = _open(tmp_path, "nu.csv", ["Time", "One"], [t, np.ones_like(t)])
    d = client.post("/api/derived", json={"op": "integral", "sources": [sigs["One"]["id"]]}).json()
    st = client.post("/api/stats", json={"signals": [d["id"]]}).json()["stats"][d["id"]]
    assert abs(st["max"] - (t[-1] - t[0])) < 1e-9
    assert d["name"] == "∫ One dt"
    client.delete(f"/api/derived/{d['name']}")
    client.delete(f"/api/runs/{run['id']}")


def test_derived_recomputes_when_source_grows(tmp_path):
    t = np.arange(100) * 0.1
    sigs, run = _open(tmp_path, "g.csv", ["Time", "V"], [t, t])
    d = client.post("/api/derived", json={"op": "mul", "sources": [sigs["V"]["id"], sigs["V"]["id"]]}).json()
    with open(tmp_path / "g.csv", "a") as f:
        f.write("10.0,10.0\n")
    client.post("/api/stream/poll")
    st = client.post("/api/stats", json={"signals": [d["id"]]}).json()["stats"][d["id"]]
    assert st["samples"] == 101 and abs(st["max"] - 100.0) < 1e-9
    client.delete(f"/api/derived/{d['name']}")
    client.delete(f"/api/runs/{run['id']}")


def test_binary_needs_overlap_and_errors_are_clear(tmp_path):
    a, ra = _open(tmp_path, "a.csv", ["Time", "A"], [np.arange(10.0), np.arange(10.0)])
    b, rb = _open(tmp_path, "b.csv", ["Time", "B"], [np.arange(20.0, 30.0), np.arange(10.0)])
    r = client.post("/api/derived", json={"op": "sub", "sources": [a["A"]["id"], b["B"]["id"]]})
    assert r.status_code == 400 and "overlap" in r.json()["detail"]
    r = client.post("/api/derived", json={"op": "lowpass", "sources": [a["A"]["id"]], "params": {"cutoff": 0}})
    assert r.status_code == 400
    client.delete(f"/api/runs/{ra['id']}"); client.delete(f"/api/runs/{rb['id']}")


def test_fft_peak_xy_and_states(tmp_path):
    t = np.arange(4000) / 1000.0
    sigs, run = _open(tmp_path, "f.csv", ["Time", "S", "C", "Gear"],
                      [t, np.sin(2 * np.pi * 50 * t), np.cos(2 * np.pi * 50 * t), np.floor(t)])
    f = client.post("/api/fft", json={"signals": [sigs["S"]["id"]]}).json()["series"][0]
    assert abs(f["peak"] - 50) < 0.5
    xy = client.post("/api/xy", json={"x": sigs["S"]["id"], "ys": [sigs["C"]["id"]]}).json()["series"][0]
    x = np.frombuffer(base64.b64decode(xy["x"]), np.float64)
    y = np.frombuffer(base64.b64decode(xy["y"]), np.float32)
    assert np.allclose(x ** 2 + y ** 2, 1, atol=1e-3)          # a circle
    st = client.post("/api/states", json={"signals": [sigs["Gear"]["id"]]}).json()["states"][sigs["Gear"]["id"]]
    assert [p[1] for p in st["points"]] == [0, 1, 2, 3] and not st["tooMany"]
    assert sigs["Gear"]["discrete"] and not sigs["S"]["discrete"]
    client.delete(f"/api/runs/{run['id']}")


# ------------------------------------------------------------------ phase 4

def _three_runs(tmp_path, shift_c=0.6, gain_c=0.9):
    t = np.arange(0, 30, 0.005)
    f = lambda tt: 20 + 10 * np.tanh(tt - 5) - 8 * np.tanh(tt - 20)
    out = []
    for name, sh, g in (("a", 0, 1), ("b", 0, 1.01), ("c", shift_c, gain_c)):
        sigs, run = _open(tmp_path, f"{name}.csv", ["Time", "Speed"], [t, g * f(t - sh)])
        out.append((sigs["Speed"]["id"], run))
    return out


def test_compare_metrics_tolerance_and_auto_shift(tmp_path):
    (a, ra), (b, rb), (c, rc) = _three_runs(tmp_path)
    r = client.post("/api/compare", json={"signals": [a, b, c], "baseline": a, "tolAbs": 1.0}).json()
    rows = {x["id"]: x for x in r["rows"]}
    assert rows[b]["pass"] and not rows[c]["pass"]
    assert rows[b]["rmsDiff"] < rows[c]["rmsDiff"]
    shift = client.post("/api/compare/shift", json={"base": a, "other": c}).json()["shift"]
    assert abs(shift + 0.6) < 0.01
    ranked = client.post("/api/compare/all", json={"runs": [ra["id"], rb["id"], rc["id"]], "baselineRun": "mean"}).json()["signals"]
    assert ranked[0]["name"] == "Speed"
    for run in (ra, rb, rc):
        client.delete(f"/api/runs/{run['id']}")


def test_reports_and_exports(tmp_path):
    (a, ra), (b, rb), (c, rc) = _three_runs(tmp_path)
    plot = {"title": "Speed", "mode": "time", "signals": [{"id": a, "label": "A", "color": "#3D8BFD"}, {"id": c, "label": "C", "color": "#E5677A"}]}
    req = {"title": "דוח", "intro": "מבוא", "summary": "סיכום", "lang": "he", "tabs": [{"name": "t", "plots": [plot]}]}
    for fmt in ("html", "pdf", "docx"):
        r = client.post("/api/report", json={**req, "format": fmt, "path": str(tmp_path / "rep")})
        assert r.status_code == 200 and os.path.getsize(r.json()["path"]) > 10_000
    html_text = open(tmp_path / "rep.html", encoding="utf-8").read()
    assert 'dir="rtl"' in html_text and "uPlot" in html_text and "sv-data" in html_text
    for fmt in ("png", "svg"):
        assert client.post("/api/export/image", json={"path": str(tmp_path / "img"), "format": fmt, "plot": plot}).status_code == 200
    e = client.post("/api/export/csv", json={"path": str(tmp_path / "d"), "signals": [{"id": a, "label": "A"}, {"id": c, "label": "C", "scale": 2}], "t0": 1, "t1": 2}).json()
    lines = open(e["path"], encoding="utf-8-sig").read().splitlines()
    assert lines[0] == "Time,A,C" and e["rows"] == 201
    for run in (ra, rb, rc):
        client.delete(f"/api/runs/{run['id']}")


def test_session_roundtrip_and_v2_import(tmp_path):
    state = {"ws": {"tabs": [], "activeTab": ""}, "runs": [], "derived": [], "ui": {"lang": "he"}}
    p = client.post("/api/session/save", json={"path": str(tmp_path / "s"), "state": state}).json()["path"]
    back = client.post("/api/session/load", json={"path": p}).json()
    assert back["state"]["ui"]["lang"] == "he" and not back["converted"]
    v2 = {"run_paths": ["/x/a.csv"], "view_state": {"layout_rows": 1, "layout_cols": 1,
          "subplots": [{"mode": "time", "assigned_signals": ["0:Speed", "-1:derived"]}]},
          "signal_settings": {"0:Speed": {"color": "#123456", "scale": 2}}, "derived_signals": {"derived": {}}}
    (tmp_path / "old.json").write_text(json.dumps(v2))
    conv = client.post("/api/session/load", json={"path": str(tmp_path / "old.json")}).json()
    sig = conv["state"]["ws"]["tabs"][0]["plots"][0]["signals"]
    assert conv["converted"] and sig == [{"sig": "r0/Speed", "color": "#123456", "width": 1.5, "scale": 2, "offset": 0, "label": None, "state": False}]
    assert conv["state"]["skippedDerived"] == 1


# ------------------------------------------------------------------ phase 5

def test_block_summary_decimation_matches_exact():
    from svcore.decimate import _minmax_range
    rng = np.random.default_rng(7)
    n = 3_000_007
    y = rng.standard_normal(n)
    y[rng.integers(0, n, 40)] = np.nan
    y[1_234_567] = 99.0
    for a, b in [(0, n), (12_345, n - 999), (n // 2, n)]:
        idx = _minmax_range(y, a, b, 3200)
        assert np.all(np.diff(idx) > 0)
        assert np.nanmax(y[idx]) == np.nanmax(y[a:b]) and np.nanmin(y[idx]) == np.nanmin(y[a:b])
    assert 1_234_567 in _minmax_range(y, 0, n, 3200)


def test_job_progress_and_cancel(tmp_path):
    import time as _t
    p = tmp_path / "big.csv"
    t = np.arange(600_000) / 1000.0
    np.savetxt(p, np.column_stack([t, np.sin(t), np.cos(t)]), delimiter=",", header="Time,A,B", comments="", fmt="%.6f")
    import svcore.csv_loader as L
    old = L.PART_BYTES
    L.PART_BYTES = 1 << 20                       # many small parts -> several progress steps
    try:
        j = client.post("/api/jobs", json={"kind": "open", "body": {"paths": [str(p)]}}).json()
        seen = set()
        while True:
            s = client.get(f"/api/jobs/{j['id']}").json()
            seen.add(round(s["progress"], 1))
            if s["state"] != "running":
                break
            _t.sleep(0.02)
        assert s["state"] == "done" and s["result"]["opened"][0]["samples"] == 600_000
        j = client.post("/api/jobs", json={"kind": "open", "body": {"paths": [str(p)]}}).json()
        client.post(f"/api/jobs/{j['id']}/cancel")
        while client.get(f"/api/jobs/{j['id']}").json()["state"] == "running":
            _t.sleep(0.02)
        assert client.get(f"/api/jobs/{j['id']}").json()["state"] in ("cancelled", "done")
    finally:
        L.PART_BYTES = old
        for rid in list(store.order):
            store.remove(rid)


def test_wide_file_single_cache_file_and_lazy_discrete(tmp_path, monkeypatch):
    import svcore.csv_loader as L
    monkeypatch.setattr(L, "CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(L, "CACHE_MIN_BYTES", 0)
    monkeypatch.setattr(L, "CACHE_ENABLED", True)
    n, k = 2000, 600
    t = np.arange(n) * 0.01
    p = tmp_path / "wide.csv"
    data = np.column_stack([t] + [np.sin(t * (i + 1)) for i in range(k - 1)] + [np.floor(t / 5)])
    np.savetxt(p, data, delimiter=",", header=",".join(["Time"] + [f"g{i // 50}.s{i}" for i in range(k)]), comments="", fmt="%.6g")
    first = L.load_csv(str(p), [])
    again = L.load_csv(str(p), [])                     # from the cache
    entry = os.listdir(tmp_path / "cache")[0]
    assert sorted(os.listdir(tmp_path / "cache" / entry)) == ["meta.json", "signals.npy", "time.npy"]
    assert all(np.array_equal(np.asarray(first.signals[s].data), np.asarray(again.signals[s].data)) for s in first.signals)
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    assert len(run["signals"]) == k and run["signals"][0]["discrete"] is None     # skipped at load
    last = run["signals"][-1]["id"]
    assert client.get("/api/discrete", params={"sig": last}).json()["discrete"] is True
    client.delete(f"/api/runs/{run['id']}")


def test_formula_signals_and_safety(tmp_path):
    t = np.arange(0, 10, 0.01)
    sigs, run = _open(tmp_path, "fx.csv", ["Time", "A", "B"], [t, np.sin(t), np.cos(t)])
    a, b = sigs["A"]["id"], sigs["B"]["id"]
    r = client.post("/api/derived", json={"op": "expr", "sources": [a, b], "params": {"expr": "2*v0 - v1 + where(v0 > 0.5, 1, 0)", "text": "2*A - B + where(A > 0.5, 1, 0)"}}).json()
    assert r["name"] == "2*A - B + where(A > 0.5, 1, 0)"
    st = client.post("/api/stats", json={"signals": [r["id"]]}).json()["stats"][r["id"]]
    expected = 2 * np.sin(t) - np.cos(t) + (np.sin(t) > 0.5)
    assert abs(st["max"] - expected.max()) < 1e-9
    for bad in ['__import__("os").system("x")', "v0.__class__", "v0[0]", "open('f')", "lambda: 1", "v7"]:
        res = client.post("/api/derived", json={"op": "expr", "sources": [a], "params": {"expr": bad}})
        assert res.status_code == 400, bad
    client.delete(f"/api/runs/{run['id']}")


def test_no_server_error_on_empty_regions(tmp_path):
    """A region with no data (e.g. cursors outside a short file) once returned NaN -> HTTP 500."""
    t = np.arange(1000) / 1000.0
    sigs, run = _open(tmp_path, "short.csv", ["Time", "S"], [t, np.sin(2 * np.pi * 50 * t)])
    sid = sigs["S"]["id"]
    for url, body in (("/api/fft", {"signals": [sid], "t0": 8, "t1": 21}),
                      ("/api/stats", {"signals": [sid], "t0": 50, "t1": 60}),
                      ("/api/data", {"signals": [sid], "x0": 50, "x1": 60, "points": 100}),
                      ("/api/cursor", {"signals": [sid], "t": 99.0})):
        r = client.post(url, json=body)
        assert r.status_code == 200, (url, r.text[:200])
    assert client.post("/api/fft", json={"signals": [sid], "t0": 8, "t1": 21}).json()["series"][0]["peak"] is None
    client.delete(f"/api/runs/{run['id']}")


def test_change_on_signals_themselves(tmp_path):
    t = np.arange(0, 10, 0.01)
    sigs, run = _open(tmp_path, "tf.csv", ["Time", "A", "B"], [t, np.sin(t), np.cos(t)])
    a, b = sigs["A"]["id"], sigs["B"]["id"]
    r = client.post("/api/transforms", json={"signals": [a, b], "transform": {"expr": "abs(x)", "scale": 10, "offset": 5, "dt": 1}})
    assert r.status_code == 200 and r.json()["previous"] == {a: None, b: None}
    st = client.post("/api/stats", json={"signals": [a]}).json()["stats"][a]
    assert abs(st["min"] - 5) < 1e-6 and abs(st["max"] - 15) < 0.01                 # 10*|sin| + 5
    assert abs(client.post("/api/cursor", json={"signals": [a], "t": 1.0 + 1.0}).json()["values"][a] - (10 * abs(np.sin(1.0)) + 5)) < 1e-6
    d = client.post("/api/derived", json={"op": "negate", "sources": [a]}).json()      # derived signals see it too
    assert abs(client.post("/api/stats", json={"signals": [d["id"]]}).json()["stats"][d["id"]]["max"] + 5) < 1e-6
    assert client.post("/api/transforms", json={"signals": [a], "transform": {"expr": "__import__('os')"}}).status_code == 400
    client.post("/api/transforms", json={"specs": r.json()["previous"]})               # undo
    assert abs(client.post("/api/stats", json={"signals": [a]}).json()["stats"][a]["max"] - 1) < 0.01
    client.delete(f"/api/derived/{d['name']}"); client.delete(f"/api/runs/{run['id']}")


def test_live_mode_detects_resaves_and_hand_typed_rows(tmp_path):
    """Live mode used to compare file SIZES only: a re-save that kept the size was missed and a
    larger re-save was read as 'appended rows'; a last row without a newline never appeared."""
    import time as _t
    p = tmp_path / "live.csv"
    def save(n, v):
        p.write_bytes(("\ufeffTime,A\r\n" + "".join(f"{i * 0.01:.2f},{v}\r\n" for i in range(n))).encode("utf-8"))
    def stat(sid): return client.post("/api/stats", json={"signals": [sid]}).json()["stats"][sid]
    save(500, "1.000")
    o = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    sid = o["signals"][0]["id"]
    _t.sleep(1.1); save(500, "2.000")                       # same size, new values
    assert client.post("/api/stream/poll").json()["changed"] == [o["id"]] and stat(sid)["max"] == 2.0
    _t.sleep(1.1); save(500, "333.000")                     # larger re-save: NOT appended rows
    client.post("/api/stream/poll")
    assert stat(sid)["max"] == 333.0 and stat(sid)["samples"] == 500
    with open(p, "a", newline="") as f:
        f.write("9.00,5.5")                                 # typed by hand, no Enter
    client.post("/api/stream/poll")                         # waits one poll (a writer may be mid-line)
    assert stat(sid)["samples"] == 500
    client.post("/api/stream/poll")
    assert stat(sid)["samples"] == 501
    client.delete(f"/api/runs/{o['id']}")


def test_live_mode_reports_unreadable_files(tmp_path, monkeypatch):
    p = tmp_path / "x.csv"
    p.write_text("Time,A\n0,1\n0.1,2\n")
    o = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    p.write_text("Time,A\n0,9\n")                           # shrank -> needs a reload ...
    def locked(rid): raise PermissionError("[Errno 13] Permission denied")
    monkeypatch.setattr(store, "reload", locked)            # ... but the file is locked
    r = client.post("/api/stream/poll").json()
    assert r["errors"] == {o["id"]: "[Errno 13] Permission denied"}   # reported, not swallowed
    monkeypatch.undo()
    client.delete(f"/api/runs/{o['id']}")


def test_restoring_formulas_is_idempotent(tmp_path):
    t = np.arange(0, 5, 0.01)
    sigs, run = _open(tmp_path, "idem.csv", ["Time", "A"], [t, np.sin(t)])
    item = {"name": "twice_A", "op": "expr", "sources": [sigs["A"]["id"]], "params": {"expr": "2*v0", "text": "2*A"}}
    client.post("/api/derived/restore", json={"items": [item]})
    r = client.post("/api/derived/restore", json={"items": [item]}).json()          # again: replaces, no duplicate
    names = [s["name"] for x in r["runs"] if x["format"] == "derived" for s in x["signals"]]
    assert names.count("twice_A") == 1 and not any(n.startswith("twice_A (") for n in names)
    client.delete("/api/derived/twice_A"); client.delete(f"/api/runs/{run['id']}")


def test_xy_returns_point_times(csv_file):
    rid = client.post("/api/files/open", json={"paths": [csv_file]}).json()["opened"][0]["id"]
    r = client.post("/api/xy", json={"x": f"{rid}/Sine", "ys": [f"{rid}/Sine"]}).json()["series"][0]
    t = np.frombuffer(base64.b64decode(r["t"]), dtype=np.float64)
    assert len(t) == r["n"] and np.all(np.diff(t) > 0)          # ascending: the timeline cuts by time


def test_object_track_aligns_roles_and_derives_course(tmp_path):
    t = np.arange(0, 60, 0.1)
    p = tmp_path / "obj.csv"
    lat = 32.0 + t * 1e-4                   # moving north
    lon = np.full_like(t, 34.8)
    with open(p, "w") as f:
        f.write("Time,Lat,Lon,Yaw\n")
        for i in range(len(t)):
            f.write(f"{t[i]:.2f},{lat[i]:.7f},{lon[i]:.7f},{(350 + i * 0.1) % 360 - 180:.3f}\n")
    rid = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]["id"]
    r = client.post("/api/object_track", json={"lat": f"{rid}/Lat", "lon": f"{rid}/Lon", "yaw": f"{rid}/Yaw"}).json()
    dec = lambda k, dt: np.frombuffer(base64.b64decode(r[k]), dtype=dt)
    assert r["n"] == len(t)
    assert np.allclose(dec("lat", np.float64)[:3], lat[:3])
    course = dec("course", np.float32)
    assert np.all(np.abs(course[5:-5]) < 1)                       # north
    assert r["units"]["yaw"] == "deg"
    assert np.all(np.isfinite(dec("yaw", np.float32)))


def test_object_track_mavlink_integer_degrees(tmp_path):
    p = tmp_path / "mav.csv"
    with open(p, "w") as f:
        f.write("Time,lat,lon\n" + "".join(f"{i},{320800000 + i * 10},{347800000}\n" for i in range(50)))
    rid = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]["id"]
    r = client.post("/api/object_track", json={"lat": f"{rid}/lat", "lon": f"{rid}/lon"}).json()
    lat = np.frombuffer(base64.b64decode(r["lat"]), dtype=np.float64)
    assert abs(lat[0] - 32.08) < 1e-6                              # 1e7 degrees -> degrees


def test_object_track_angle_units_by_the_object(tmp_path):
    # pitch / roll of a few degrees: degrees, because the yaw of the same object is clearly degrees
    p = tmp_path / "att.csv"
    with open(p, "w") as f:
        f.write("Time,Lat,Lon,Yaw,Pitch,Roll,Yrad\n")
        for i in range(200):
            f.write(f"{i * 0.1},{32 + i * 1e-5},{34.8},{(i * 2) % 360 - 180},{5 * np.sin(i / 10):.3f},{3 * np.cos(i / 10):.3f},{np.sin(i / 20):.3f}\n")
    rid = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]["id"]
    r = client.post("/api/object_track", json={"lat": f"{rid}/Lat", "lon": f"{rid}/Lon", "yaw": f"{rid}/Yaw",
                                               "pitch": f"{rid}/Pitch", "roll": f"{rid}/Roll"}).json()
    assert r["units"] == {"yaw": "deg", "pitch": "deg", "roll": "deg"}
    r = client.post("/api/object_track", json={"lat": f"{rid}/Lat", "lon": f"{rid}/Lon", "yaw": f"{rid}/Yrad"}).json()
    assert r["units"]["yaw"] == "rad"                              # all within +-pi: radians


def test_object_track_from_a_quaternion(tmp_path):
    # heading 90 deg (east), pitch 10 deg, roll -20 deg as a w-x-y-z quaternion (ZYX order, like PX4)
    import math
    yaw, pitch, roll = math.radians(90), math.radians(10), math.radians(-20)
    cy, sy, cp, sp, cr, sr = math.cos(yaw / 2), math.sin(yaw / 2), math.cos(pitch / 2), math.sin(pitch / 2), math.cos(roll / 2), math.sin(roll / 2)
    w, x, y, z = cr * cp * cy + sr * sp * sy, sr * cp * cy - cr * sp * sy, cr * sp * cy + sr * cp * sy, cr * cp * sy - sr * sp * cy
    p = tmp_path / "quat.csv"
    with open(p, "w") as f:
        f.write("Time,lat,lon,q[0],q[1],q[2],q[3]\n" + "".join(f"{i},{32 + i * 1e-5},34.8,{w},{x},{y},{z}\n" for i in range(50)))
    rid = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]["id"]
    r = client.post("/api/object_track", json={"lat": f"{rid}/lat", "lon": f"{rid}/lon", "qw": f"{rid}/q[0]", "qx": f"{rid}/q[1]",
                                               "qy": f"{rid}/q[2]", "qz": f"{rid}/q[3]"}).json()
    dec = lambda k: np.frombuffer(base64.b64decode(r[k]), dtype=np.float32)
    assert r["units"] == {"yaw": "rad", "pitch": "rad", "roll": "rad"}
    assert abs(math.degrees(dec("yaw")[10]) - 90) < 0.1
    assert abs(math.degrees(dec("pitch")[10]) - 10) < 0.1
    assert abs(math.degrees(dec("roll")[10]) + 20) < 0.1


def test_save_image_png_only(tmp_path):
    png = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20).decode()
    r = client.post("/api/save/image", json={"path": str(tmp_path / "view.png"), "data": "data:image/png;base64," + png}).json()
    assert open(r["path"], "rb").read().startswith(b"\x89PNG")
    assert client.post("/api/save/image", json={"path": str(tmp_path / "x.png"), "data": base64.b64encode(b"hello").decode()}).status_code == 400


def test_decimation_keeps_context_outside_the_window():
    """The view window gets full detail; the rest stays as coarse context (not clipped).
    A window past the end of the data must still return the trace, so panning never blanks."""
    x = np.linspace(0.0, 60.0, 60_000)
    y = np.sin(x)
    xd, yd, dec = decimate_for_view(x, y, (100.0, 200.0), 400)
    assert dec and 0 < len(xd) < 2000 and xd[-1] == x[-1]
    xd, yd, _ = decimate_for_view(x, y, (10.0, 11.0), 400)
    inside = (xd >= 10.0) & (xd <= 11.0)
    assert inside.sum() >= 300                     # the window keeps its full budget
    assert xd[0] == x[0] and xd[-1] == x[-1]       # coarse context on both sides
    assert np.all(np.diff(xd) >= 0)                # still sorted for the chart


def test_fs_list_and_complete(tmp_path):
    (tmp_path / "flights").mkdir()
    (tmp_path / "Flight_notes").mkdir()
    (tmp_path / "flight_012.csv").write_text("t,a\n0,1\n")
    (tmp_path / "flight_013.bin").write_bytes(b"x")
    (tmp_path / "flyer.pdf").write_text("x")
    (tmp_path / ".hidden").mkdir()
    r = client.post("/api/fs/list", json={"path": str(tmp_path)}).json()
    names = [e["name"] for e in r["entries"]]
    assert names[:2] == ["Flight_notes", "flights"] and "flyer.pdf" not in names and ".hidden" not in names
    assert all("mtime" in e for e in r["entries"])

    sep = os.sep
    r = client.post("/api/fs/complete", json={"text": str(tmp_path) + sep + "fli"}).json()
    got = [i["name"] for i in r["items"]]
    assert got == ["Flight_notes", "flights", "flight_012.csv", "flight_013.bin"]   # folders first, any case
    assert r["common"] == str(tmp_path) + sep + "Flight"                            # what Tab fills in
    assert r["items"][0]["path"].endswith(sep)                                      # a folder continues

    r = client.post("/api/fs/complete", json={"text": str(tmp_path) + sep + "flights"}).json()
    assert r["common"] == str(tmp_path / "flights") + sep                           # single match: whole path

    r = client.post("/api/fs/complete", json={"text": "fl", "base": str(tmp_path), "exts": [".csv"]}).json()
    assert [i["name"] for i in r["items"]] == ["Flight_notes", "flights", "flight_012.csv"]   # relative + exts

    r = client.post("/api/fs/complete", json={"text": str(tmp_path) + sep + "nothing" + sep + "x"}).json()
    assert r["items"] == []
    r = client.post("/api/fs/complete", json={"text": str(tmp_path) + sep + ".h"}).json()
    assert [i["name"] for i in r["items"]] == [".hidden"]                          # dot only when typed


def test_examples_listed_with_existing_files_only(monkeypatch, tmp_path):
    from app.routes import fs as m
    r = client.get("/api/examples").json()
    assert [x["id"] for x in r] == ["flight", "car"] and len(r[1]["paths"]) == 3
    assert all(os.path.isfile(p) for x in r for p in x["paths"])
    (tmp_path / "uav_flight.csv").write_text("t,a\n0,1\n")
    monkeypatch.setattr(m, "EXAMPLES", str(tmp_path))
    assert [x["id"] for x in client.get("/api/examples").json()] == ["flight"]     # car files missing


def test_every_example_ships_in_every_package():
    """F10: the welcome screen's examples exist in the repo and every package copies examples/ to
    where routes/fs.py looks (beside backend/). The Windows portable build once left them out and
    the welcome line silently disappeared."""
    from app.routes import fs as m
    root = os.path.join(os.path.dirname(__file__), "..", "..")
    for ex in m.EXAMPLE_SETS:
        for f in ex["files"]:
            assert os.path.isfile(os.path.join(root, "examples", f)), f
    # from source, fs.py's default is <the folder that holds backend>/examples
    default = os.path.join(os.path.dirname(m.__file__), "..", "..", "..", "examples")
    assert os.path.samefile(default, os.path.join(root, "examples"))
    with open(os.path.join(root, "tools", "make_windows_portable.py"), encoding="utf-8") as f:
        portable = f.read()
    assert 'os.path.join(app, "backend")' in portable and 'os.path.join(app, "examples")' in portable
    with open(os.path.join(root, "SignalViewer.spec"), encoding="utf-8") as f:
        assert '"examples"), "examples")' in f.read()          # PyInstaller; run.py sets SV_EXAMPLES


def test_formula_angles_and_quaternions():
    """T17: deg/rad/wrap and quaternion -> Euler (aerospace order, degrees)."""
    from svcore.formula import evaluate
    t = np.linspace(0, 1, 5)
    ang = np.radians([10.0, -30.0, 45.0])            # roll, pitch, yaw
    cr, sr = np.cos(ang[0] / 2), np.sin(ang[0] / 2)
    cp, sp = np.cos(ang[1] / 2), np.sin(ang[1] / 2)
    cy, sy = np.cos(ang[2] / 2), np.sin(ang[2] / 2)
    q = [cr * cp * cy + sr * sp * sy, sr * cp * cy - cr * sp * sy, cr * sp * cy + sr * cp * sy, cr * cp * sy - sr * sp * cy]
    sigs = [(t, np.full_like(t, c)) for c in q]
    interp = lambda tb, ts, ys: np.interp(tb, ts, ys)
    for fn, want in (("q2roll", 10.0), ("q2pitch", -30.0), ("q2yaw", 45.0)):
        _, y = evaluate(f"{fn}(v0, v1, v2, v3)", sigs, interp)
        assert np.allclose(y, want, atol=1e-9), fn
    _, y = evaluate("wrap180(v0)", [(t, np.array([190.0, -190.0, 180.0, 0.0, 359.0]))], interp)
    assert np.allclose(y, [-170.0, 170.0, -180.0, 0.0, -1.0])
    _, y = evaluate("wrap360(deg(rad(v0)))", [(t, np.array([-10.0, 370.0, 0.0, 90.0, 720.0]))], interp)
    assert np.allclose(y, [350.0, 10.0, 0.0, 90.0, 0.0])


def _quat(roll, pitch, yaw):
    """Hamilton w-first quaternion, body -> NED, from aerospace angles in degrees."""
    r, p, y = np.radians([roll, pitch, yaw]) / 2
    cr, sr, cp, sp, cy, sy = np.cos(r), np.sin(r), np.cos(p), np.sin(p), np.cos(y), np.sin(y)
    return [cr * cp * cy + sr * sp * sy, sr * cp * cy - cr * sp * sy,
            cr * sp * cy + sr * cp * sy, cr * cp * sy - sr * sp * cy]


def _frame3(sigs, **params):
    from svcore.analysis import compute
    return np.vstack([compute("frame", sigs, {**params, "axis": k})[1] for k in range(3)])


def test_frame_ned_enu():
    """T17b: NED <-> ENU swaps north / east and flips the vertical, both ways."""
    t = np.linspace(0, 1, 4)
    ned = [(t, np.full_like(t, c)) for c in (1.0, 2.0, 3.0)]
    assert np.allclose(_frame3(ned, src="ned", dst="enu"), [[2.0] * 4, [1.0] * 4, [-3.0] * 4])
    enu = [(t, np.full_like(t, c)) for c in (2.0, 1.0, -3.0)]
    assert np.allclose(_frame3(enu, src="enu", dst="ned"), [[1.0] * 4, [2.0] * 4, [3.0] * 4])


def test_frame_body_euler_and_quaternion_agree():
    """T17b: body -> NED with roll / pitch / yaw (deg or rad) equals the quaternion result, the
    known cases come out right, and NED -> body undoes body -> NED."""
    t = np.linspace(0, 1, 3)
    c = lambda v: (t, np.full_like(t, float(v)))
    fwd = [c(1), c(0), c(0)]
    # nose east: forward is east; nose up 30 deg: forward is north and up (negative down)
    assert np.allclose(_frame3(fwd + [c(0), c(0), c(90)], src="body", dst="ned")[:, 0], [0, 1, 0], atol=1e-12)
    assert np.allclose(_frame3(fwd + [c(0), c(30), c(0)], src="body", dst="ned")[:, 0],
                       [np.cos(np.radians(30)), 0, -0.5], atol=1e-12)
    assert np.allclose(_frame3(fwd + [c(0), c(0), c(90)], src="body", dst="enu")[:, 0], [1, 0, 0], atol=1e-12)
    rng = np.random.default_rng(7)
    for _ in range(20):
        att = rng.uniform([-180, -89, -180], [180, 89, 180])
        vec = [c(x) for x in rng.normal(size=3)]
        eul = _frame3(vec + [c(a) for a in att], src="body", dst="ned")
        rad = _frame3(vec + [c(a) for a in np.radians(att)], src="body", dst="ned", unit="rad")
        qua = _frame3(vec + [c(x * 1.01) for x in _quat(*att)], src="body", dst="ned", att="quat")  # not unit length
        assert np.allclose(eul, rad) and np.allclose(eul, qua)
        back = _frame3([c(x) for x in eul[:, 0]] + [c(a) for a in att], src="ned", dst="body")
        assert np.allclose(back[:, 0], [v[1][0] for v in vec])


def test_frame_attitude_interpolation():
    """T17b: the attitude is aligned onto the vector's samples without the classic traps - yaw
    crossing +-180 deg, and a quaternion that flips sign (q and -q are the same attitude)."""
    tv, ta = np.array([0.25, 0.5, 0.75]), np.array([0.0, 1.0])
    fwd = [(tv, np.ones(3)), (tv, np.zeros(3)), (tv, np.zeros(3))]
    zero = (ta, np.zeros(2))
    yaw = (ta, np.array([179.0, -179.0]))                 # halfway is 180 (south), not 0 (north)
    assert np.allclose(_frame3(fwd + [zero, zero, yaw], src="body", dst="ned")[:, 1], [-1, 0, 0], atol=1e-9)
    q = np.array(_quat(0, 0, 90))
    flip = [(ta, np.array([q[k], -q[k]])) for k in range(4)]
    assert np.allclose(_frame3(fwd + flip, src="body", dst="ned", att="quat")[:, 1], [0, 1, 0], atol=1e-9)


def test_frame_lla_ecef():
    """T17b: LLA <-> ECEF on the WGS84 ellipsoid - known points, and back to within a millimetre."""
    from svcore.analysis import WGS84_A, WGS84_E2, ecef_to_lla, lla_to_ecef
    assert np.allclose(lla_to_ecef(0, 0, 0), [WGS84_A, 0, 0])
    assert np.allclose(lla_to_ecef(90, 0, 0), [0, 0, WGS84_A * np.sqrt(1 - WGS84_E2)], atol=1e-6)
    assert np.allclose(lla_to_ecef(0, 90, 100), [0, WGS84_A + 100, 0], atol=1e-6)
    rng = np.random.default_rng(3)
    lat, lon, alt = rng.uniform(-90, 90, 200), rng.uniform(-180, 180, 200), rng.uniform(-400, 40000, 200)
    back = ecef_to_lla(*lla_to_ecef(lat, lon, alt))
    assert np.allclose(back[0], lat, atol=1e-9) and np.allclose(back[1], lon, atol=1e-9)
    assert np.allclose(back[2], alt, atol=1e-3)


def test_frame_position_and_vector_differ():
    """T17b: a position moves the origin, a vector only rotates. The same ECEF numbers give 0 as a
    position at the reference point, but a 6378 km 'down' arrow as a vector."""
    from svcore.analysis import WGS84_A, lla_to_ecef
    t = np.linspace(0, 1, 3)
    c = lambda v: (t, np.full_like(t, float(v)))
    at_ref = [c(WGS84_A), c(0), c(0)]                   # lat 0, lon 0, alt 0
    pos = _frame3(at_ref, qty="position", src="ecef", dst="ned", ref="first")
    assert np.allclose(pos, 0, atol=1e-6)
    vec = _frame3(at_ref + [c(0), c(0)], src="ecef", dst="ned")
    assert np.allclose(vec[:, 0], [0, 0, -WGS84_A])
    # 100 m up and ~111 m north of a manual reference in Tel Aviv
    ref = dict(ref="manual", ref_lat=32.08, ref_lon=34.78, ref_alt=30.0)
    lla = [c(32.081), c(34.78), c(130)]
    ned = _frame3(lla, qty="position", src="lla", dst="ned", **ref)[:, 0]
    assert abs(ned[0] - 110.9) < 0.5 and abs(ned[1]) < 1e-6 and abs(ned[2] + 100) < 0.01
    enu = _frame3(lla, qty="position", src="lla", dst="enu", **ref)[:, 0]
    assert np.allclose(enu, [ned[1], ned[0], -ned[2]])
    back = _frame3([c(x) for x in ned], qty="position", src="ned", dst="lla", **ref)[:, 0]
    assert np.allclose(back, [32.081, 34.78, 130], atol=1e-8)
    ecef = _frame3(lla, qty="position", src="lla", dst="ecef")[:, 0]
    assert np.allclose(ecef, lla_to_ecef(32.081, 34.78, 130))
    # ECEF velocity <-> NED velocity uses the vehicle's lat / lon, and comes back
    rng = np.random.default_rng(5)
    for _ in range(10):
        la, lo, vn = rng.uniform(-80, 80), rng.uniform(-180, 180), rng.normal(size=3)
        e = _frame3([c(x) for x in vn] + [c(la), c(lo)], src="ned", dst="ecef")[:, 0]
        assert np.isclose(np.linalg.norm(e), np.linalg.norm(vn))          # a rotation keeps the length
        assert np.allclose(_frame3([c(x) for x in e] + [c(la), c(lo)], src="ecef", dst="ned")[:, 0], vn)
    north = _frame3([c(1), c(0), c(0), c(0), c(0)], src="ned", dst="ecef")[:, 0]
    assert np.allclose(north, [0, 0, 1])                # at lat 0 / lon 0, north is ECEF +Z


def test_frame_longitude_across_antimeridian():
    """T17b: a track crossing +-180 deg is interpolated across the line, not around the globe."""
    from svcore.analysis import WGS84_A
    tp, tv = np.array([0.0, 1.0]), np.array([0.0, 0.5, 1.0])
    lla = [(tv, np.zeros(3)), (tp, np.array([179.9, -179.9])), (tv, np.zeros(3))]
    x, y, _ = _frame3(lla, qty="position", src="lla", dst="ecef")
    assert np.isclose(x[1], -WGS84_A) and abs(y[1]) < 1e-6           # on the line, not at lon 0
    east = _frame3(lla, qty="position", src="lla", dst="ned", ref="first")[1]
    # the reference point is the origin: 0 up to rounding (macOS / Windows math libraries give ~1e-13)
    assert abs(east[0]) < 1e-6 and abs(east[1] - east[2] / 2) < 1.0 and 22000 < east[2] < 22500


def test_frame_errors_say_what_to_do():
    from svcore.analysis import OpError, compute
    t = np.linspace(0, 1, 3)
    three = [(t, t), (t, t), (t, t)]
    cases = [
        ({"qty": "position", "src": "body", "dst": "ned"}, "not a place"),
        ({"qty": "vector", "src": "lla", "dst": "ned"}, "choose Position"),
        ({"qty": "position", "src": "ned", "dst": "lla", "ref": "first"}, "reference point"),
        ({"qty": "position", "src": "ned", "dst": "lla", "ref": "manual", "ref_lat": 95, "ref_lon": 0, "ref_alt": 0}, "-90 and 90"),
        ({"src": "ned", "dst": "ecef"}, "latitude and longitude"),
        ({"qty": "speed", "src": "ned", "dst": "enu"}, "vector"),
    ]
    for params, text in cases:
        with pytest.raises(OpError, match=text):
            compute("frame", three, params)


def test_frame_api_three_signals(tmp_path):
    """T17b: one conversion = three ordinary derived signals (axis 0 / 1 / 2) with their own names;
    they restore like any formula, and wrong input says what is needed."""
    p = tmp_path / "vel.csv"
    with open(p, "w") as f:
        f.write("Time,vn,ve,vd\n" + "".join(f"{i * 0.1},{1 + i},{2 * i},{-1}\n" for i in range(20)))
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    src = [f"{run['id']}/{n}" for n in ("vn", "ve", "vd")]
    ids = []
    for k in range(3):
        r = client.post("/api/derived", json={"op": "frame", "sources": src,
                                               "params": {"src": "ned", "dst": "enu", "axis": k}})
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
    assert ids == ["derived/E NED→ENU(vn, ve, vd)", "derived/N NED→ENU(vn, ve, vd)", "derived/U NED→ENU(vn, ve, vd)"]
    up = client.post("/api/data", json={"signals": [ids[2]], "points": 100}).json()["series"][0]
    assert np.allclose(np.frombuffer(base64.b64decode(up["y"]), dtype=np.float32), 1.0)
    items = [{"op": "frame", "sources": src, "params": {"src": "ned", "dst": "enu", "axis": k},
              "name": i.split("/", 1)[1]} for k, i in enumerate(ids)]
    r = client.post("/api/derived/restore", json={"items": items}).json()
    assert not r["failed"] and len(r["restored"]) == 3
    bad = client.post("/api/derived", json={"op": "frame", "sources": src, "params": {"src": "body", "dst": "ned"}})
    assert bad.status_code == 400 and "6 signals" in bad.json()["detail"]
    same = client.post("/api/derived", json={"op": "frame", "sources": src, "params": {"src": "ned", "dst": "ned"}})
    assert same.status_code == 400
    r = client.post("/api/derived", json={"op": "frame", "sources": src,
                                           "params": {"qty": "position", "src": "ned", "dst": "enu", "axis": 2}})
    assert r.status_code == 200 and r.json()["id"] == "derived/U pos NED→ENU(vn, ve, vd)"
    ids.append(r.json()["id"])
    for i in ids:
        store.remove_derived(i.split("/", 1)[1])
    store.remove(run["id"])


def test_segments_of_a_state_signal():
    """T10: flight modes as segments; totals per mode add up to the recording."""
    ex = os.path.join(os.path.dirname(__file__), "..", "..", "examples", "uav_flight.csv")
    run = client.post("/api/files/open", json={"paths": [ex]}).json()["opened"][0]
    sid = f"{run['id']}/FlightMode"
    r = client.post("/api/segments", json={"signal": sid}).json()
    segs = r["segments"]
    assert segs and r["total"] == len(segs)
    assert all(a["t1"] == b["t0"] for a, b in zip(segs, segs[1:]))           # back to back
    assert [s["label"] for s in segs][:2] == ["IDLE", "TAKEOFF"]
    total = sum(x["duration"] for x in r["byValue"])
    assert abs(total - (segs[-1]["t1"] - segs[0]["t0"])) < 1e-6
    # the refactor kept "statistics while state = X" the same: AUTO time = its segments' time
    auto = next(x for x in r["byValue"] if x["label"] == "AUTO")
    st = client.post("/api/stats", json={"signals": [sid], "t0": None, "t1": None,
                                         "state": sid, "stateValues": [auto["value"]]}).json()
    assert st["stats"][sid]["samples"] > 0
    assert client.post("/api/segments", json={"signal": "nope/x"}).status_code == 404
    client.delete(f"/api/runs/{run['id']}")


def test_welch_psd_finds_the_tone_and_is_smoother(tmp_path):
    """T16: Welch PSD through the API: peak at the tone; far less scatter than one FFT."""
    t = np.arange(0, 60, 0.002)                                   # 500 Hz, 60 s
    rng = np.random.default_rng(1)
    y = np.sin(2 * np.pi * 37.0 * t) + rng.normal(0, 2.0, len(t))
    p = tmp_path / "vib.csv"
    np.savetxt(p, np.column_stack([t, y]), delimiter=",", header="Time,AccZ", comments="")
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    sid = f"{run['id']}/AccZ"
    dec = lambda s, k, dt: np.frombuffer(base64.b64decode(s[k]), dtype=dt)
    w = client.post("/api/fft", json={"signals": [sid], "method": "welch"}).json()["series"][0]
    f = client.post("/api/fft", json={"signals": [sid]}).json()["series"][0]
    assert abs(w["peak"] - 37.0) < 1.0 and abs(f["peak"] - 37.0) < 0.5
    wy, fy = dec(w, "y", np.float32), dec(f, "y", np.float32)
    noise = lambda a: np.std(np.log10(a[len(a) // 2:] + 1e-30))      # scatter of the noise floor
    assert noise(wy) < noise(fy) / 3
    client.delete(f"/api/runs/{run['id']}")


def test_limits_pass_fail():
    """T13: spans outside, time outside, worst excess; scale/offset/time offset as plotted."""
    from svcore.limits import check_limits
    t = np.arange(0, 10, 0.1)
    y = np.zeros_like(t); y[20:30] = 5.0; y[60:65] = -3.0
    r = check_limits(t, y, -2.0, 4.0)
    assert not r["pass"] and r["violations"] == 2
    assert np.allclose(r["spans"], [[2.0, 3.0], [6.0, 6.5]])
    assert abs(r["timeOutside"] - 1.5) < 1e-9 and r["worst"] == 1.0 and r["first"] == 2.0
    assert check_limits(t, y, None, 6.0)["pass"] and check_limits(t, y, None, None)["pass"]
    assert check_limits(t, y, -2.0, 4.0, 4.0, 5.5)["pass"]                       # window inside
    p = os.path.join(os.path.dirname(__file__), "..", "..", "examples", "run_012.csv")
    run = client.post("/api/files/open", json={"paths": [p]}).json()["opened"][0]
    sid = f"{run['id']}/Speed [km/h]"
    q = lambda **k: client.post("/api/limits", json={"items": [{"id": sid, **k}]}).json()["results"][sid]
    assert q(hi=80)["pass"] and not q(hi=70)["pass"]
    assert not q(hi=80, scale=2)["pass"]                                         # checked as plotted
    a, b = q(hi=70)["first"], q(hi=70, tOffset=5)["first"]
    assert abs(b - a - 5) < 1e-9
    client.delete(f"/api/runs/{run['id']}")


def test_report_limits_column_in_html_pdf_docx(tmp_path):
    """T13: a signal with limits gets a Limits column (pass / fail) in all three formats;
    a report without limits keeps its old table."""
    p = os.path.join(os.path.dirname(__file__), "..", "..", "examples", "run_012.csv")
    run = client.post("/api/files/open", json={"paths": [p]}).json()["opened"][0]
    sig = lambda n, **k: {"id": f"{run['id']}/{n}", "label": n, "color": "#3D8BFD", "width": 1.5, "scale": 1, "offset": 0, **k}
    plot = {"mode": "time", "title": "", "signals": [sig("Speed [km/h]", limits={"lo": None, "hi": 70}),
                                                      sig("RPM", limits={"lo": 0, "hi": 5000})]}
    body = lambda fmt, plots: {"path": str(tmp_path / f"r.{fmt}"), "format": fmt, "tabs": [{"name": "T", "plots": plots}]}
    client.post("/api/report", json=body("html", [plot]))
    h = (tmp_path / "r.html").read_text(encoding="utf-8")
    assert "<th>Limits</th>" in h and 'class="lim-fail"' in h and 'class="lim-pass"' in h
    for fmt in ("pdf", "docx"):
        assert client.post("/api/report", json=body(fmt, [plot])).status_code == 200
    import docx
    cells = [c.text for tb in docx.Document(str(tmp_path / "r.docx")).tables for row in tb.rows for c in row.cells]
    assert "Limits" in cells and any(c.startswith("Fail") for c in cells) and any(c.startswith("Pass") for c in cells)
    plain = {"mode": "time", "title": "", "signals": [sig("RPM")]}
    client.post("/api/report", json=body("html", [plain]))
    assert "Limits" not in (tmp_path / "r.html").read_text(encoding="utf-8")
    client.delete(f"/api/runs/{run['id']}")


def test_envelope_bands_and_outlier(tmp_path):
    """T14: 12 runs of one signal (one far off) -> ordered bands over the overlap, one outlier."""
    rng = np.random.default_rng(3)
    ids, runs = [], []
    for k in range(12):
        t = np.arange(0.0 + 0.1 * k, 20.0, 0.01)
        y = np.sin(t) * (1 + 0.05 * rng.normal()) + (3.0 if k == 7 else 0.0)
        p = tmp_path / f"mc_{k:02d}.csv"
        np.savetxt(p, np.column_stack([t, y]), delimiter=",", header="Time,Alt", comments="")
        r = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
        runs.append(r["id"]); ids.append({"id": f"{r['id']}/Alt"})
    out = client.post("/api/envelope", json={"signals": ids, "points": 500}).json()
    dec = lambda k, dt=np.float32: np.frombuffer(base64.b64decode(out[k]), dtype=dt)
    t, mn, lo, med, hi, mx = dec("t", np.float64), dec("min"), dec("lo"), dec("median"), dec("hi"), dec("max")
    assert len(t) == 500 and abs(t[0] - 1.1) < 1e-9 and abs(t[-1] - 19.99) < 1e-6      # common overlap
    assert np.all(mn <= lo + 1e-6) and np.all(lo <= med + 1e-6) and np.all(med <= hi + 1e-6) and np.all(hi <= mx + 1e-6)
    assert [r["outlier"] for r in out["runs"]].count(True) == 1 and out["runs"][7]["outlier"]
    assert client.post("/api/envelope", json={"signals": ids[:1]}).json()["empty"]
    for r in runs:
        client.delete(f"/api/runs/{r}")


def test_envelope_no_outlier_in_an_ordinary_spread():
    """T14: a normal random spread has no outliers; two runs far off are both found."""
    from svcore.envelope import envelope
    rng = np.random.default_rng(11)
    t = np.arange(0, 10, 0.01)
    runs = [(t, np.sin(t) * (1 + 0.05 * rng.normal())) for _ in range(30)]
    assert not any(envelope(runs, points=400)["outlier"])
    runs[3] = (t, runs[3][1] + 5); runs[17] = (t, runs[17][1] - 4)
    out = envelope(runs, points=400)["outlier"]
    assert [i for i, o in enumerate(out) if o] == [3, 17]


def test_report_envelope_plot_in_html_pdf_docx(tmp_path):
    """T14b: an Envelope plot in a report - the bands' data in HTML, a summary caption in every
    format (Hebrew too), and PDF / Word build."""
    ids, runs = [], []
    for k in range(4):
        t = np.arange(0, 10, 0.01)
        p = tmp_path / f"r{k}.csv"
        np.savetxt(p, np.column_stack([t, np.sin(t) + (6 if k == 2 else 0.02 * k)]), delimiter=",", header="Time,V", comments="")
        r = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
        runs.append(r["id"]); ids.append({"id": f"{r['id']}/V", "label": f"run{k}", "color": "#3D8BFD", "scale": 1, "offset": 0})
    plot = {"mode": "env", "title": "", "caption": "", "signals": ids}
    body = lambda fmt, lang="en": {"path": str(tmp_path / f"e.{fmt}"), "format": fmt, "lang": lang, "tabs": [{"name": "T", "plots": [plot]}]}
    assert client.post("/api/report", json=body("html")).status_code == 200
    h = (tmp_path / "e.html").read_text(encoding="utf-8")
    assert '"env"' in h and "4 runs" in h and "run2" in h                       # bands + caption + outlier named
    assert client.post("/api/report", json=body("html", "he")).status_code == 200
    assert "4 ריצות" in (tmp_path / "e.html").read_text(encoding="utf-8")
    for fmt in ("pdf", "docx"):
        assert client.post("/api/report", json=body(fmt)).status_code == 200
    import docx
    text = "\n".join(par.text for par in docx.Document(str(tmp_path / "e.docx")).paragraphs)
    assert "4 runs" in text and "run2" in text
    from pypdf import PdfReader                                                  # the plot page is there
    assert "4 runs" in "".join(pg.extract_text() or "" for pg in PdfReader(str(tmp_path / "e.pdf")).pages)
    for r in runs:
        client.delete(f"/api/runs/{r}")


def test_xyz_track_aligns_and_thins():
    """T17c: Y and Z on X's samples over the time all three share; thinned keeps the last sample."""
    from svcore.xyz import xyz_track
    tx, ty = np.linspace(0, 10, 1001), np.linspace(2, 12, 51)
    r = xyz_track((tx, tx * 2), (ty, ty * 3), (tx, -tx))
    assert r["t"][0] == 2.0 and r["t"][-1] == 10.0
    assert np.allclose(r["y"], r["t"] * 3) and np.allclose(r["z"], -r["t"]) and np.allclose(r["x"], r["t"] * 2)
    thin = xyz_track((tx, tx), (tx, tx), (tx, tx), points=100)
    assert len(thin["t"]) <= 101 and thin["t"][-1] == 10.0
    assert xyz_track((tx, tx), (tx + 20, tx), (tx, tx)) is None                # no common time


def test_xyz_api_and_report(tmp_path):
    """T17c: /api/xyz returns the tracks; an XYZ plot is a real 3D chart in HTML (image), PDF and Word."""
    p = tmp_path / "traj.csv"
    t = np.arange(0, 20, 0.05)
    np.savetxt(p, np.column_stack([t, 50 * np.cos(t / 3), 50 * np.sin(t / 3), -2 * t]), delimiter=",",
               header="Time,pn,pe,pd", comments="")
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    ids = [f"{run['id']}/{n}" for n in ("pn", "pe", "pd")]
    r = client.post("/api/xyz", json={"tracks": [{"x": ids[0], "y": ids[1], "z": ids[2]},
                                                 {"x": "nope/x", "y": ids[1], "z": ids[2]}]}).json()["tracks"]
    assert r[0]["n"] == len(t) and r[1]["missing"]
    z = np.frombuffer(base64.b64decode(r[0]["z"]), dtype=np.float32)
    assert np.allclose(z, -2 * t, atol=1e-4)
    sigs = [{"id": i, "label": i.split("/")[1], "color": "#3D8BFD", "scale": 1, "offset": 0} for i in ids]
    plot = {"mode": "xyz", "title": "Trajectory", "caption": "", "signals": sigs, "xyzFlipZ": True}
    body = lambda fmt: {"path": str(tmp_path / f"x.{fmt}"), "format": fmt, "lang": "en", "tabs": [{"name": "T", "plots": [plot]}]}
    assert client.post("/api/report", json=body("html")).status_code == 200
    h = (tmp_path / "x.html").read_text(encoding="utf-8")
    assert 'class="specimg xyzimg"' in h and "data:image/png;base64," in h
    for fmt in ("pdf", "docx"):
        assert client.post("/api/report", json=body(fmt)).status_code == 200
    from pypdf import PdfReader
    assert "Trajectory" in "".join(pg.extract_text() or "" for pg in PdfReader(str(tmp_path / "x.pdf")).pages)
    import docx
    assert len(docx.Document(str(tmp_path / "x.docx")).inline_shapes) >= 1
    store.remove(run["id"])


def test_frame_group_nested_names(tmp_path):
    """The three axes of one conversion form a nested group: named "<group>.<axis>" with
    params.group, kept through the listing, the preview and a restore (sessions / autosave)."""
    p = tmp_path / "pos.csv"
    with open(p, "w") as f:
        f.write("Time,n,e,d\n" + "".join(f"{i * 0.1},{i},{2 * i},{-i}\n" for i in range(20)))
    run = client.post("/api/files/open", json={"paths": [str(p)]}).json()["opened"][0]
    src = [f"{run['id']}/{n}" for n in ("n", "e", "d")]
    base = {"qty": "position", "src": "ned", "dst": "enu", "group": "enu_Pos_ned"}
    items = [{"op": "frame", "sources": src, "params": {**base, "axis": k}, "name": f"enu_Pos_ned.{a}"}
             for k, a in enumerate(("E", "N", "U"))]
    for it in items:
        r = client.post("/api/derived", json=it)
        assert r.status_code == 200 and r.json()["id"] == f"derived/{it['name']}"
    derived = next(r for r in client.get("/api/runs").json() if r["format"] == "derived")
    mine = [s for s in derived["signals"] if s["params"].get("group") == "enu_Pos_ned"]
    assert [s["name"] for s in mine] == ["enu_Pos_ned.E", "enu_Pos_ned.N", "enu_Pos_ned.U"]
    pv = client.post("/api/derived/preview", json={"op": "frame", "sources": src, "params": {**base, "axis": 2}}).json()
    assert np.allclose(np.frombuffer(base64.b64decode(pv["y"]), dtype=np.float32), [i for i in range(20)])  # U = -D
    r = client.post("/api/derived/restore", json={"items": items}).json()
    assert not r["failed"] and r["restored"] == [it["name"] for it in items]       # same names: no "(2)"
    for it in items:
        store.remove_derived(it["name"])
    store.remove(run["id"])


def test_frame_reference_from_signals():
    """T17c: a local track goes onto the map with the origin taken from other signals' first valid
    sample (the vehicle's GPS) - which may cover a different time than the track."""
    t = np.linspace(10, 20, 11)
    c = lambda v, tt=t: (tt, np.full_like(tt, float(v)))
    tg = np.linspace(0, 30, 31)
    lat = (tg, np.where(tg < 2, np.nan, 32.08))                        # GPS without a fix at first
    ned = [c(0), c(0), c(-100)]                                         # 100 m above the origin
    lla = _frame3(ned + [lat, c(34.78, tg), c(30, tg)], qty="position", src="ned", dst="lla", ref="signals")
    assert np.allclose(lla[:, 0], [32.08, 34.78, 130], atol=1e-7) and len(lla[0]) == 11
    from svcore.analysis import OpError, compute
    with pytest.raises(OpError, match="6 signals"):
        compute("frame", ned, {"qty": "position", "src": "ned", "dst": "lla", "ref": "signals"})
