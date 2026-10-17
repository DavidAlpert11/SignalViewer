"""
Reports (phase 4): interactive standalone HTML, PDF, Word, PNG/SVG images.

All statistics are computed on full-resolution data. The HTML report embeds the
data as min/max resolution levels so zooming in shows more detail with no server.
"""

import logging
import base64
import html
import io
import json
import os
from datetime import datetime
from typing import Dict, List, Optional

import numpy as np

from svcore.limits import check_limits
from svcore.envelope import envelope
from svcore.analysis import fft_display, fft_for, transitions, xy_pair
from svcore.decimate import decimate_for_view, minmax_decimate
from svcore.ops import compute_signal_stats

log = logging.getLogger("signalviewer").getChild(__name__)

ASSETS = os.path.join(os.path.dirname(__file__), "report_assets")
LEVELS = (4000, 40000, 150000)         # min/max points per level for long time signals
RAW_LIMIT = 150000                     # signals up to this length are embedded raw

STRINGS = {
    "en": {"generated": "Generated", "summary": "Summary", "notes": "Notes", "toggle": "Click to show or hide", "theme": "Light / dark",
           "min": "Min", "max": "Max", "mean": "Mean", "rms": "RMS", "std": "Std", "samples": "Samples",
           "signal": "Signal", "hint": "Drag to zoom · wheel to zoom · double-click to reset · click to place the cursor",
           "page": "Page", "plot": "Plot", "peak": "Peak", "visible": "Statistics for the shown time range",
           "limits": "Limits", "pass": "Pass", "fail": "Fail",
           "envNote": "{n} runs · median, {p}–{q} % band, min–max", "envFar": "Far from the rest:"},
    "he": {"generated": "נוצר", "summary": "סיכום", "notes": "הערות", "toggle": "לחיצה להצגה או הסתרה", "theme": "בהיר / כהה",
           "min": "מינ׳", "max": "מקס׳", "mean": "ממוצע", "rms": "RMS", "std": "סטיית תקן", "samples": "דגימות",
           "signal": "סיגנל", "hint": "גרירה לזום · גלגלת לזום · לחיצה כפולה לאיפוס · לחיצה למיקום הסמן",
           "page": "עמוד", "plot": "גרף", "peak": "שיא", "visible": "סטטיסטיקות לטווח הזמן המוצג",
           "limits": "גבולות", "pass": "עבר", "fail": "נכשל",
           "envNote": "{n} ריצות · חציון, רצועת {p}–{q} %, מינ׳–מקס׳", "envFar": "רחוקות מהשאר:"},
}


def _b64(a, dtype) -> str:
    return base64.b64encode(np.ascontiguousarray(a, dtype=dtype).tobytes()).decode("ascii")


def _f(v):
    return None if v is None or not np.isfinite(v) else float(v)


# =============================================================================
# DATA FOR ONE PLOT (shared by every format)
# =============================================================================

def plot_data(store, spec: Dict, for_html: bool, points: int = 3000) -> Dict:
    """Resolve a plot spec from the UI into display data + full-resolution statistics."""
    mode = spec.get("mode") or "time"
    rng = spec.get("range")
    out = {"title": spec.get("title") or "", "caption": spec.get("caption") or "", "mode": mode,
           "range": rng, "log": bool(spec.get("fftLog", True)), "signals": [], "notes": spec.get("notes") or []}
    if mode == "spec":
        return _spec_data(store, spec, out)
    if mode == "env":
        return _env_data(store, spec, out, for_html)
    if mode == "xyz":
        return _xyz_data(store, spec, out)
    if mode == "xy":
        xs = spec.get("xSignal")
        if not xs:
            return out
        tx, x = store.signal(xs)
        out["xLabel"] = spec.get("xLabel") or xs.split("/", 1)[-1]
    for s in spec.get("signals", []):
        sid = s["id"]
        try:
            t, y = store.signal(sid)
        except KeyError:
            continue
        scale, offset = float(s.get("scale", 1) or 1), float(s.get("offset", 0) or 0)
        t_off = float(s.get("tOffset", 0) or 0)
        if t_off:
            t = np.asarray(t, dtype=np.float64) + t_off
        ys = np.asarray(y, dtype=np.float64) * scale + offset if (scale != 1 or offset != 0) else y
        sig = {"label": s.get("label") or sid.split("/", 1)[-1], "color": s.get("color", "#3D8BFD"),
               "width": s.get("width", 1.5), "state": bool(s.get("state"))}
        labels = store.labels(sid) if hasattr(store, "labels") and not sid.startswith("derived/") else None
        if labels and scale == 1 and offset == 0:
            st = None                                # text values: no numeric statistics
        else:
            st = compute_signal_stats(t, ys, rng[0] if rng else None, rng[1] if rng else None) if mode == "time" else \
                compute_signal_stats(t, ys)
        sig["stats"] = {k: _f(st.get(k)) for k in ("min", "max", "mean", "rms", "std", "samples")} if st else {}
        lim = s.get("limits") or {}
        if mode == "time" and (lim.get("lo") is not None or lim.get("hi") is not None):
            sig["limits"] = check_limits(t, ys, lim.get("lo"), lim.get("hi"),
                                         rng[0] if rng else None, rng[1] if rng else None)
        if mode == "time":
            if for_html:
                if len(t) <= RAW_LIMIT:
                    levels = [(t, ys)]
                else:
                    levels = [minmax_decimate(t, ys, n) for n in LEVELS if n < len(t)]
                sig["levels"] = [{"x": _b64(a, np.float64), "y": _b64(b, np.float32), "n": int(len(a))} for a, b in levels]
            else:
                xd, yd, _ = decimate_for_view(t, ys, rng, points)
                sig["x"], sig["y"] = np.asarray(xd, dtype=np.float64), np.asarray(yd, dtype=np.float64)
            if sig["state"]:
                pts, many, _end = transitions(t, ys, cap=200)
                def txt(v):
                    if labels and float(v).is_integer() and int(v) in labels:
                        return labels[int(v)]
                    return int(v) if float(v).is_integer() else v
                sig["states"] = [] if many else [[p[0], txt(p[1])] for p in pts if p[1] is not None]
        elif mode == "fft":
            welch = spec.get("fftMethod") == "welch"
            f, m = fft_for(sid, store.version(sid), t, y, spec.get("fftWindow") or "hanning", None, None,
                           "welch" if welch else "fft")
            fx, fy, peak = fft_display(f, m * (scale * scale if welch else abs(scale)), 4000 if for_html else points)
            sig["peak"] = _f(peak)
            sig["x"], sig["y"] = np.asarray(fx, dtype=np.float64), np.asarray(fy, dtype=np.float64)
        else:
            a, b = xy_pair(tx, x, t, ys, spec.get("xyAlign") or "linear", points=16000 if for_html else 6000)
            sig["x"], sig["y"] = a, b
        if for_html and "x" in sig:
            sig["x"], sig["y"] = _b64(sig["x"], np.float64), _b64(sig["y"], np.float32)
        out["signals"].append(sig)
    return out


def _env_data(store, spec: Dict, out: Dict, for_html: bool) -> Dict:
    """T14b: an Envelope plot - the same computation as on screen (svcore.envelope)."""
    series, labels = [], []
    for s in spec.get("signals") or []:
        try:
            t, y = store.signal(s["id"])
        except KeyError:
            continue
        series.append((np.asarray(t, dtype=np.float64) + float(s.get("tOffset", 0) or 0),
                       np.asarray(y, dtype=np.float64) * float(s.get("scale", 1) or 1) + float(s.get("offset", 0) or 0)))
        labels.append(s.get("label") or s["id"].split("/", 1)[-1])
    rng = spec.get("range")
    r = envelope(series, rng[0] if rng else None, rng[1] if rng else None, 2000)
    sigs = spec.get("signals") or []
    out["envColor"] = (sigs[0].get("color") if sigs else None) or "#3D8BFD"
    if r is None:
        return out
    keys = ("min", "max", "lo", "hi", "median")
    if for_html:
        out["env"] = {"t": _b64(r["t"], np.float64), **{k: _b64(r[k], np.float32) for k in keys}}
    else:
        out["env"] = {"t": r["t"], **{k: r[k] for k in keys}}
    out["envRuns"] = len(series)
    out["envOutliers"] = [lab for lab, o in zip(labels, r["outlier"]) if o]
    out["envPct"] = r["pct"]
    # the summary goes with the caption: every format (HTML, PDF, Word) shows captions
    S = STRINGS.get(spec.get("lang") or "en", STRINGS["en"])
    note = S["envNote"].format(n=len(series), p=f"{r['pct']:g}", q=f"{100 - r['pct']:g}")
    if out["envOutliers"]:
        note += " · " + S["envFar"] + " " + ", ".join(out["envOutliers"][:8])
    out["caption"] = (out["caption"] + "\n" if out["caption"] else "") + note
    return out


def _spec_data(store, spec: Dict, out: Dict) -> Dict:
    """Spectrogram plot: the image of its first signal over the plot's time window."""
    from svcore.analysis import spectrogram
    sigs = spec.get("signals") or []
    if not sigs:
        return out
    s = sigs[0]
    try:
        t, y = store.signal(s["id"])
    except KeyError:
        return out
    t_off = float(s.get("tOffset", 0) or 0)
    rng = out["range"]
    try:
        r = spectrogram(t, y, (rng[0] - t_off) if rng else None, (rng[1] - t_off) if rng else None,
                        int(spec.get("specN") or 1024), 1100, 320, spec.get("specFmax"))
    except ValueError:
        return out
    step = float(r["f"][1] - r["f"][0]) if len(r["f"]) > 1 else r["df"]
    cw = float(r["t"][1] - r["t"][0]) if len(r["t"]) > 1 else r["dt"]
    out["spec"] = {"img": r["img"], "extent": [float(r["t"][0]) - cw / 2 + t_off, float(r["t"][-1]) + cw / 2 + t_off,
                                             max(0.0, float(r["f"][0]) - step / 2), float(r["f"][-1]) + step / 2],
                   "db": [r["dbLo"], r["dbHi"]]}
    out["signals"].append({"label": s.get("label") or s["id"].split("/", 1)[-1], "color": s.get("color", "#3D8BFD"),
                           "width": 1.5, "state": False, "stats": {}})
    return out


def _xyz_data(store, spec: Dict, out: Dict) -> Dict:
    """T17c: a plain 3D plot - the plot's signals in threes (X, Y, Z), each three one track, with
    scale / offset as plotted. Drawn by matplotlib in every format (HTML gets it as an image)."""
    from svcore.xyz import xyz_track
    sigs = spec.get("signals") or []
    out["xyz"], out["xyzEqual"], out["xyzFlipZ"] = [], bool(spec.get("xyzEqual", True)), bool(spec.get("xyzFlipZ"))
    for k in range(0, len(sigs) - len(sigs) % 3, 3):
        trio = sigs[k:k + 3]
        try:
            r = xyz_track(*(store.signal(s["id"]) for s in trio), spec.get("xyAlign") or "linear", 6000)
        except KeyError:
            continue
        if r is None:
            continue
        sc = [(float(s.get("scale", 1) or 1), float(s.get("offset", 0) or 0)) for s in trio]
        labels = [s.get("label") or s["id"].split("/", 1)[-1] for s in trio]
        out["xyz"].append({"x": r["x"] * sc[0][0] + sc[0][1], "y": r["y"] * sc[1][0] + sc[1][1], "z": r["z"] * sc[2][0] + sc[2][1],
                           "color": trio[0].get("color", "#3D8BFD"), "labels": labels})
        out["signals"].append({"label": " / ".join(labels), "color": trio[0].get("color", "#3D8BFD"), "width": 1.5,
                               "state": False, "stats": {}})
    return out


def _draw_xyz(ax, pdata: Dict, th: Dict):
    """A 3D axes in place of the 2D one the caller made (same position on the page)."""
    fig, pos = ax.figure, ax.get_position()
    ax.remove()
    ax = fig.add_axes(pos, projection="3d")
    ax.set_facecolor(th["bg"])
    for a in (ax.xaxis, ax.yaxis, ax.zaxis):                # panes and grid in the report theme (dark too)
        a.set_pane_color(th["bg"])
        a._axinfo["grid"]["color"] = th["grid"]
    tracks = pdata.get("xyz") or []
    for tr in tracks:
        ax.plot(tr["x"], tr["y"], tr["z"], color=tr["color"], linewidth=1.2, label=_bidi(" / ".join(tr["labels"])))
        ok = np.isfinite(tr["x"]) & np.isfinite(tr["y"]) & np.isfinite(tr["z"])
        if ok.any():
            i = int(np.flatnonzero(ok)[0])
            ax.scatter([tr["x"][i]], [tr["y"][i]], [tr["z"][i]], color=tr["color"], s=12)
    if tracks:
        lab = tracks[0]["labels"]
        for setter, text in ((ax.set_xlabel, lab[0]), (ax.set_ylabel, lab[1]), (ax.set_zlabel, lab[2])):
            setter(_bidi(text), color=th["axis"], fontsize=7)
        if pdata.get("xyzEqual"):
            try:
                ax.set_box_aspect((1, 1, 1))
                allv = [np.concatenate([tr[k][np.isfinite(tr[k])] for tr in tracks]) for k in ("x", "y", "z")]
                if all(len(v) for v in allv):
                    c = [(v.min() + v.max()) / 2 for v in allv]
                    h = max(max(v.max() - v.min() for v in allv) / 2, 1e-9)
                    ax.set_xlim(c[0] - h, c[0] + h); ax.set_ylim(c[1] - h, c[1] + h); ax.set_zlim(c[2] - h, c[2] + h)
            except (AttributeError, TypeError, ValueError):     # older matplotlib: keep its own scaling
                pass
        if pdata.get("xyzFlipZ"):
            ax.invert_zaxis()
        if len(tracks) > 1:
            leg = ax.legend(fontsize=7, loc="upper right", facecolor=th["bg"], edgecolor=th["grid"])
            for t in leg.get_texts():
                t.set_color(th["fg"])
    ax.tick_params(colors=th["axis"], labelsize=6.5)
    return ax


def _spec_png(pdata: Dict, theme: str) -> str:
    fig = render_figure(pdata, theme, size=(7.0, 5.0) if pdata.get("mode") == "xyz" else (9.0, 3.4), dpi=110)
    return base64.b64encode(figure_bytes(fig, "png")).decode()


def _notes_html(notes: List[Dict], S: Dict) -> str:
    if not notes:
        return ""
    rows = "".join('<li><span class="nt">%s s</span> %s</li>' % (html.escape(_num(n.get("t"))), html.escape(n.get("text", "")))
                   for n in sorted(notes, key=lambda n: n.get("t") or 0))
    return '<h2>%s</h2><ul class="notes">%s</ul>' % (S["notes"], rows)


def _num(v) -> str:
    try:
        return f"{float(v):.6g}"
    except (TypeError, ValueError):
        return ""


# =============================================================================
# HTML
# =============================================================================

CSS = """
:root,[data-theme=dark]{--bg:#151b23;--panel:#1c2430;--line:#2d3a4b;--ink:#dde4ec;--ink-strong:#f4f7fa;--muted:#8e9cad;
--amber:#f2a93b;--amber-ink:#1b1305;--plot:#121820;--grid:#222c39;--axis:#8e9cad}
[data-theme=light]{--bg:#eef1f4;--panel:#fff;--line:#d5dce4;--ink:#1f2a36;--ink-strong:#0f1822;--muted:#5a6979;
--amber:#945d0b;--amber-ink:#fff;--plot:#fff;--grid:#e6ebf0;--axis:#5a6979}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 "IBM Plex Sans","IBM Plex Sans Hebrew","Segoe UI",system-ui,sans-serif}
main{max-width:1100px;margin:0 auto;padding:32px 24px 64px}
header{display:flex;align-items:flex-start;gap:16px;margin-bottom:20px}
h1{margin:0;font-size:26px;font-weight:600;color:var(--ink-strong);letter-spacing:-.01em;flex:1}
.meta{color:var(--muted);font-size:13px;margin-top:4px}
button{font:inherit;color:inherit;cursor:pointer}
#theme{border:1px solid var(--line);background:var(--panel);border-radius:6px;padding:6px 10px}
.text{white-space:pre-wrap;max-width:75ch;margin:0 0 20px}
nav{display:flex;gap:4px;border-bottom:1px solid var(--line);margin:24px 0 16px}
.tabbtn{border:0;background:none;padding:8px 14px;color:var(--muted);border-bottom:2px solid transparent;margin-bottom:-1px}
.tabbtn[aria-selected=true]{color:var(--ink-strong);border-bottom-color:var(--ink-strong)}
.plot{background:var(--plot);border:1px solid var(--line);border-radius:10px;padding:14px 14px 10px;margin:0 0 22px;break-inside:avoid}
.plot h2{margin:0 0 6px;font-size:16px;font-weight:600;color:var(--ink-strong)}
.caption{color:var(--muted);margin:2px 0 8px;white-space:pre-wrap}
.specimg{width:100%;height:auto;border-radius:6px}.specimg.xyzimg{max-width:720px}
.notes{padding-inline-start:18px}.notes .nt{color:var(--muted);font-variant-numeric:tabular-nums;margin-inline-end:6px}
.legend{display:flex;flex-wrap:wrap;gap:6px;margin-bottom:6px}
.chip{display:inline-flex;align-items:center;gap:6px;border:1px solid var(--line);background:none;border-radius:12px;padding:1px 9px;font-size:12.5px}
.chip.off{opacity:.45}.chip .sw{width:12px;height:3px;border-radius:2px}.chip .nm{unicode-bidi:plaintext}
.chip .val{color:var(--amber);font-variant-numeric:tabular-nums;unicode-bidi:plaintext}
.canvas{direction:ltr;width:100%}
table{width:100%;border-collapse:collapse;margin-top:8px;font-size:12.5px;font-variant-numeric:tabular-nums}
th,td{padding:4px 8px;border-top:1px solid var(--line);text-align:end;white-space:nowrap}
th{color:var(--muted);font-weight:500}td:first-child,th:first-child{text-align:start;unicode-bidi:plaintext}
.sw2{display:inline-block;width:9px;height:9px;border-radius:2px;margin-inline-end:6px;vertical-align:middle}
.lim-pass{color:#1f8a4c;font-weight:600}.lim-fail{color:#c0392b;font-weight:600}
.hint{color:var(--muted);font-size:12px;margin:0 0 14px}
.u-select{background:rgba(242,169,59,.15)!important}
footer{color:var(--muted);font-size:12px;margin-top:40px;border-top:1px solid var(--line);padding-top:12px}
@media print{#theme,.hint,nav{display:none}.tabpane{display:block!important}body{background:#fff}}
"""


_FONT_FILES = [
    # (family, weight, file, unicode-range) - embedded so a report looks the same on any machine, offline
    ("IBM Plex Sans", 400, "ibm-plex-sans-latin-400-normal.woff2", "U+0000-00FF,U+0131,U+0152-0153,U+02BB-02BC,U+02C6,U+02DA,U+02DC,U+2000-206F,U+2074,U+20AC,U+2122,U+2191,U+2193,U+2212,U+2215,U+FEFF,U+FFFD"),
    ("IBM Plex Sans", 600, "ibm-plex-sans-latin-600-normal.woff2", "U+0000-00FF,U+0131,U+0152-0153,U+02BB-02BC,U+02C6,U+02DA,U+02DC,U+2000-206F,U+2074,U+20AC,U+2122,U+2191,U+2193,U+2212,U+2215,U+FEFF,U+FFFD"),
    ("IBM Plex Sans Hebrew", 400, "ibm-plex-sans-hebrew-hebrew-400-normal.woff2", "U+0590-05FF,U+200C-2010,U+20AA,U+25CC,U+FB1D-FB4F"),
    ("IBM Plex Sans Hebrew", 600, "ibm-plex-sans-hebrew-hebrew-600-normal.woff2", "U+0590-05FF,U+200C-2010,U+20AA,U+25CC,U+FB1D-FB4F"),
]
_font_css_cache: Optional[str] = None


def _fonts_css() -> str:
    global _font_css_cache
    if _font_css_cache is None:
        parts = []
        for family, weight, name, rng in _FONT_FILES:
            path = os.path.join(ASSETS, "fonts", name)
            if not os.path.isfile(path):
                continue
            with open(path, "rb") as f:
                data = base64.b64encode(f.read()).decode("ascii")
            parts.append("@font-face{font-family:'%s';font-style:normal;font-weight:%d;font-display:swap;"
                         "src:url(data:font/woff2;base64,%s) format('woff2');unicode-range:%s}" % (family, weight, data, rng))
        _font_css_cache = "".join(parts)
    return _font_css_cache


def _limit_text(s: Dict, S: Dict) -> str:
    """T13: the Limits cell of a report table - the same words in HTML, PDF and Word."""
    r = s.get("limits")
    if not r:
        return ""
    rng = " … ".join("" if v is None else f"{v:.6g}" for v in (r["lo"], r["hi"]))
    if r["pass"]:
        return f"{S['pass']} ({rng})"
    return f"{S['fail']} ({rng}): {r['violations']}× · {r['timeOutside']:.3g} s"


def _stats_table(sigs: List[Dict], S: Dict) -> str:
    rows = []
    for s in sigs:
        st = s.get("stats") or {}
        cells = "".join(f"<td>{'' if st.get(k) is None else f'{st[k]:.6g}'}</td>" for k in ("min", "max", "mean", "rms", "std"))
        samples = "" if st.get("samples") is None else f"{int(st['samples']):,}"
        lim = ""
        if any(x.get("limits") for x in sigs):
            r = s.get("limits")
            cls = "" if not r else (' class="lim-pass"' if r["pass"] else ' class="lim-fail"')
            lim = f"<td{cls}>{html.escape(_limit_text(s, S))}</td>"
        rows.append(f'<tr><td><span class="sw2" style="background:{html.escape(s["color"])}"></span>{html.escape(s["label"])}</td>{cells}<td>{samples}</td>{lim}</tr>')
    if not rows:
        return ""
    head = "".join(f"<th>{S[k]}</th>" for k in ("min", "max", "mean", "rms", "std", "samples"))
    if any(x.get("limits") for x in sigs):
        head += f"<th>{S['limits']}</th>"
    return f'<table><thead><tr><th>{S["signal"]}</th>{head}</tr></thead><tbody>{"".join(rows)}</tbody></table>'


def _count_plots(req: Dict) -> int:
    return max(1, sum(len(t.get("plots", [])) for t in req.get("tabs", [])))


def build_html(store, req: Dict, progress=None) -> str:
    lang = req.get("lang", "en")
    S = STRINGS.get(lang, STRINGS["en"])
    tabs, done, total = [], 0, _count_plots(req)
    for tab in req.get("tabs", []):
        plots = []
        for p in tab.get("plots", []):
            plots.append(plot_data(store, {**p, "notes": req.get("notes") or [], "lang": req.get("lang", "en")}, for_html=True))
            done += 1
            if progress:
                progress(0.05 + 0.9 * done / total, "plots")
        tabs.append({"name": tab.get("name", ""), "plots": plots})
    stats_html = {}
    for ti, tab in enumerate(tabs):
        for pi, p in enumerate(tab["plots"]):
            stats_html[(ti, pi)] = _stats_table(p["signals"], S) if p["mode"] not in ("xy", "spec", "xyz") else ""
            for s in p["signals"]:
                s.pop("stats", None)
    payload_tabs = [{"name": t["name"], "plots": [{k: v for k, v in p.items() if k not in ("spec", "xyz")} for p in t["plots"]]} for t in tabs]
    payload = {"lang": lang, "strings": {"toggle": S["toggle"]}, "tabs": payload_tabs}
    nav_items = []
    for i, t in enumerate(tabs):
        label = html.escape(t["name"] or str(i + 1))
        nav_items.append('<button class="tabbtn" aria-selected="%s">%s</button>' % (str(i == 0).lower(), label))
    nav = "" if len(tabs) < 2 else "<nav>" + "".join(nav_items) + "</nav>"
    panes = []
    for ti, t in enumerate(tabs):
        items = []
        for pi, p in enumerate(t["plots"]):
            heading = html.escape(p["title"] or "%s %d" % (S["plot"], pi + 1))
            caption = '<p class="caption">%s</p>' % html.escape(p["caption"]) if p["caption"] else ""
            if p["mode"] in ("spec", "xyz"):
                img = '<img class="specimg%s" alt="%s" src="data:image/png;base64,%s">' % (
                    " xyzimg" if p["mode"] == "xyz" else "", heading, _spec_png(p, req.get("theme", "light"))) \
                    if p.get("spec") or p.get("xyz") else ""
                items.append('<section class="plot" data-p="%d"><h2>%s</h2>%s%s</section>' % (pi, heading, caption, img))
                continue
            items.append('<section class="plot" data-p="%d"><h2>%s</h2>%s<div class="legend"></div>'
                         '<div class="canvas"></div>%s</section>' % (pi, heading, caption, stats_html[(ti, pi)]))
        panes.append(f'<div class="tabpane"{"" if ti == 0 else " hidden"}>{"".join(items)}</div>')
    def read(n):
        with open(os.path.join(ASSETS, n), encoding="utf-8") as fh:
            return fh.read()
    title = html.escape(req.get("title") or "Signal Viewer report")
    data_json = json.dumps(payload, separators=(",", ":")).replace("</", "<\\/")
    intro = '<p class="text">%s</p>' % html.escape(req["intro"]) if req.get("intro") else ""
    summary = '<h2>%s</h2><p class="text">%s</p>' % (S["summary"], html.escape(req["summary"])) if req.get("summary") else ""
    summary += _notes_html(req.get("notes") or [], S)
    return TEMPLATE.format(
        lang=lang, dir="rtl" if lang == "he" else "ltr", theme="dark" if req.get("theme") == "dark" else "light",
        title=title, uplot_css=read("uPlot.min.css"), css=_fonts_css() + CSS, generated=S["generated"],
        date=datetime.now().strftime("%Y-%m-%d %H:%M"), theme_label=S["theme"], intro=intro, hint=S["hint"],
        nav=nav, panes="".join(panes), summary=summary, data=data_json,
        uplot_js=read("uPlot.iife.min.js"), viewer_js=read("viewer.js"))


TEMPLATE = """<!doctype html>
<html lang="{lang}" dir="{dir}" data-theme="{theme}">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><style>{uplot_css}</style><style>{css}</style></head>
<body><main>
<header><div style="flex:1"><h1>{title}</h1><div class="meta">{generated} {date}</div></div>
<button id="theme" title="{theme_label}">&#9680;</button></header>
{intro}
<p class="hint">{hint}</p>
{nav}{panes}
{summary}
<footer>Signal Viewer</footer>
</main>
<script type="application/json" id="sv-data">{data}</script>
<script>{uplot_js}</script>
<script>{viewer_js}</script>
</body></html>"""


# =============================================================================
# STATIC FIGURES (PDF / Word / PNG / SVG) with matplotlib
# =============================================================================

def _bidi(text: str) -> str:
    if not text:
        return ""
    try:
        from bidi.algorithm import get_display
    except ImportError:                 # optional: Hebrew then shows in logical order
        return text
    try:
        return get_display(text)
    except Exception:                   # odd text: draw it as it is rather than fail the report
        log.debug("bidi failed for %r", text[:40], exc_info=True)
        return text


FIG_THEME = {
    "light": {"bg": "#ffffff", "fg": "#1f2a36", "grid": "#e6ebf0", "axis": "#5a6979"},
    "dark": {"bg": "#121820", "fg": "#dde4ec", "grid": "#222c39", "axis": "#8e9cad"},
}


def draw_axes(ax, pdata: Dict, theme: str = "light"):
    th = FIG_THEME.get(theme, FIG_THEME["light"])
    ax.set_facecolor(th["bg"])
    for sp in ax.spines.values():
        sp.set_color(th["grid"])
    ax.tick_params(colors=th["axis"], labelsize=8)
    ax.grid(True, color=th["grid"], linewidth=0.8)
    ax.set_axisbelow(True)
    mode = pdata["mode"]
    if mode == "xyz":
        _draw_xyz(ax, pdata, th)
        return
    if mode == "env":
        e = pdata.get("env")
        if e is not None:
            c = pdata.get("envColor") or "#3D8BFD"
            ax.fill_between(e["t"], e["min"], e["max"], color=c, alpha=0.15, linewidth=0)
            ax.fill_between(e["t"], e["lo"], e["hi"], color=c, alpha=0.35, linewidth=0)
            ax.plot(e["t"], e["median"], color=c, linewidth=1.4)
            if pdata.get("range"):
                ax.set_xlim(*pdata["range"])
        return
    if mode == "spec":
        sp = pdata.get("spec")
        if sp:
            ax.imshow(sp["img"], origin="lower", aspect="auto", extent=sp["extent"], cmap="viridis", interpolation="nearest")
            ax.set_ylabel("Hz", color=th["axis"], fontsize=8)
            ax.grid(False)
            if pdata["signals"]:
                ax.set_title(_bidi(pdata["signals"][0]["label"]), fontsize=8, color=th["axis"], loc="left")
        return
    for s in pdata["signals"]:
        x, y = s.get("x"), s.get("y")
        if x is None or not len(x):
            continue
        kw = dict(color=s["color"], linewidth=float(s.get("width") or 1.5) * 0.8, label=_bidi(s["label"]))
        if mode == "time" and s.get("state"):
            ax.step(x, y, where="post", **kw)
            for i, (tv, val) in enumerate(s.get("states") or []):
                if i:
                    ax.axvline(tv, color=s["color"], alpha=0.45, linewidth=0.8)
                ax.annotate(str(val), (tv, 1), xycoords=("data", "axes fraction"), xytext=(2, -12), textcoords="offset points",
                            fontsize=7, color="white", bbox=dict(boxstyle="square,pad=0.2", fc=s["color"], ec="none"))
        else:
            ax.plot(x, y, **kw)
    if mode == "fft":
        ax.set_xlabel("Hz", color=th["axis"], fontsize=8)
        if pdata.get("log"):
            ax.set_yscale("log")
    elif mode == "xy":
        ax.set_xlabel(_bidi(pdata.get("xLabel", "")), color=th["axis"], fontsize=8)
    if mode == "time":
        for n in pdata.get("notes") or []:                   # the user's notes on time plots
            try:
                tv = float(n.get("t"))
            except (TypeError, ValueError):
                continue
            ax.axvline(tv, color="#3D8BFD", linestyle="--", linewidth=0.8, alpha=0.8)
            ax.annotate(_bidi(str(n.get("text", ""))[:40]), (tv, 1), xycoords=("data", "axes fraction"), xytext=(3, -10),
                        textcoords="offset points", fontsize=6.5, color="#3D8BFD")
    if mode == "time" and pdata.get("range"):
        ax.set_xlim(*pdata["range"])
    if pdata["signals"]:
        leg = ax.legend(fontsize=7.5, loc="upper right", frameon=True, facecolor=th["bg"], edgecolor=th["grid"])
        for t in leg.get_texts():
            t.set_color(th["fg"])


def render_figure(pdata: Dict, theme: str = "light", size=(9.0, 3.6), dpi: int = 150):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=size, dpi=dpi)
    fig.patch.set_facecolor(FIG_THEME.get(theme, FIG_THEME["light"])["bg"])
    draw_axes(ax, pdata, theme)
    if pdata.get("mode") != "xyz":                          # a 3D axes does not take part in tight_layout
        fig.tight_layout()
    return fig


def figure_bytes(fig, fmt: str) -> bytes:
    import matplotlib.pyplot as plt
    buf = io.BytesIO()
    fig.savefig(buf, format=fmt, facecolor=fig.get_facecolor())
    plt.close(fig)
    return buf.getvalue()


def _collect(store, req: Dict, progress=None) -> List[Dict]:
    items, done, total = [], 0, _count_plots(req)
    for tab in req.get("tabs", []):
        for p in tab.get("plots", []):
            pdata = plot_data(store, {**p, "notes": req.get("notes") or [], "lang": req.get("lang", "en")}, for_html=False)
            if pdata["signals"] or pdata.get("env") is not None:      # an Envelope plot keeps its data in "env"
                items.append(pdata)
            done += 1
            if progress:
                progress(0.05 + 0.6 * done / total, "plots")
    return items


def build_pdf(store, req: Dict, path: str, progress=None):
    """A4 pages: a title page (title, date, introduction), then one page per plot with
    heading, chart, caption and full-resolution statistics, then the summary."""
    import textwrap
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    lang = req.get("lang", "en")
    S = STRINGS.get(lang, STRINGS["en"])
    rtl = lang == "he"
    ha, X = ("right", 0.9) if rtl else ("left", 0.1)
    ink, muted = "#0f1822", "#5a6979"
    items = _collect(store, req, progress)
    title = req.get("title") or "Signal Viewer report"
    total = 1 + len(items) + (1 if req.get("summary") or req.get("notes") else 0)

    def page(n):
        fig = plt.figure(figsize=(8.27, 11.69))
        fig.text(X, 0.965, _bidi(title), ha=ha, fontsize=8, color=muted)
        fig.text(1 - X if rtl else 0.9, 0.965, _bidi("%s %d / %d" % (S["page"], n, total)),
                 ha="left" if rtl else "right", fontsize=8, color=muted)
        return fig

    def paragraphs(fig, text, y, size=10.5, width=88, color=ink):
        for para in (text or "").split("\n"):
            for line in textwrap.wrap(para, width) or [""]:
                fig.text(X, y, _bidi(line), ha=ha, va="top", fontsize=size, color=color)
                y -= 0.0175 * size / 10
        return y

    n = 1
    with PdfPages(path) as pdf:
        fig = page(n)
        fig.text(X, 0.86, _bidi(title), ha=ha, va="top", fontsize=22, weight="bold", color=ink)
        fig.text(X, 0.815, _bidi("%s %s" % (S["generated"], datetime.now().strftime("%Y-%m-%d %H:%M"))), ha=ha, fontsize=9, color=muted)
        paragraphs(fig, req.get("intro"), 0.77)
        pdf.savefig(fig); plt.close(fig)

        for i, p in enumerate(items):
            if progress:
                progress(0.65 + 0.33 * i / max(1, len(items)), "pages")
            n += 1
            fig = page(n)
            fig.text(X, 0.92, _bidi(p["title"] or "%s %d" % (S["plot"], i + 1)), ha=ha, va="top", fontsize=14, weight="bold", color=ink)
            ax = fig.add_axes([0.1, 0.52, 0.82, 0.36])
            draw_axes(ax, p, "light")
            y = 0.47
            if p["caption"]:
                y = paragraphs(fig, p["caption"], y, size=9.5, color=muted) - 0.01
            rows = [s for s in p["signals"] if s.get("stats")]
            if rows and p["mode"] != "xy":
                cols = [S["signal"], S["min"], S["max"], S["mean"], S["rms"], S["std"]]
                cells = [[_bidi(s["label"])] + ["" if s["stats"].get(k) is None else "%.6g" % s["stats"][k]
                          for k in ("min", "max", "mean", "rms", "std")] for s in rows]
                if any(s.get("limits") for s in rows):
                    cols = cols + [S["limits"]]
                    cells = [c + [_bidi(_limit_text(s, S))] for c, s in zip(cells, rows)]
                if rtl:  # columns read right-to-left
                    cols, cells = cols[::-1], [r[::-1] for r in cells]
                h = min(0.04 * (len(rows) + 1), y - 0.08)
                tax = fig.add_axes([0.1, y - h, 0.82, h]); tax.axis("off")
                tbl = tax.table(cellText=cells, colLabels=[_bidi(c) for c in cols], loc="upper center", cellLoc="center")
                tbl.auto_set_font_size(False); tbl.set_fontsize(8.5); tbl.scale(1, 1.25)
                for (r, c), cell in tbl.get_celld().items():
                    cell.set_edgecolor("#d5dce4")
                    if r == 0:
                        cell.set_facecolor("#eef2f5"); cell.set_text_props(weight="bold", color=ink)
            pdf.savefig(fig); plt.close(fig)

        notes = sorted(req.get("notes") or [], key=lambda n: n.get("t") or 0)
        if req.get("summary") or notes:
            n += 1
            fig = page(n)
            y = 0.9
            if req.get("summary"):
                fig.text(X, y, _bidi(S["summary"]), ha=ha, va="top", fontsize=16, weight="bold", color=ink)
                y = paragraphs(fig, req["summary"], y - 0.04) - 0.03
            if notes:
                fig.text(X, y, _bidi(S["notes"]), ha=ha, va="top", fontsize=14, weight="bold", color=ink)
                paragraphs(fig, "\n".join(f"{_num(nt.get('t'))} s  {nt.get('text', '')}" for nt in notes), y - 0.04)
            pdf.savefig(fig); plt.close(fig)


def build_docx(store, req: Dict, path: str, progress=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt
    lang = req.get("lang", "en")
    S = STRINGS.get(lang, STRINGS["en"])
    rtl = lang == "he"
    doc = Document()

    def rtlize(par):
        if rtl:
            ppr = par._p.get_or_add_pPr()
            b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); ppr.append(b)
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in par.runs:
                rpr = r._r.get_or_add_rPr()
                rt = OxmlElement("w:rtl"); rt.set(qn("w:val"), "1"); rpr.append(rt)
        return par

    rtlize(doc.add_heading(req.get("title") or "Signal Viewer report", 0))
    rtlize(doc.add_paragraph(f"{S['generated']} {datetime.now().strftime('%Y-%m-%d %H:%M')}"))
    for para in (req.get("intro") or "").split("\n"):
        if para.strip():
            rtlize(doc.add_paragraph(para))
    items = _collect(store, req, progress)
    for i, p in enumerate(items):
        if progress:
            progress(0.65 + 0.33 * i / max(1, len(items)), "pages")
        rtlize(doc.add_heading(p["title"] or f"{S['plot']} {i + 1}", 1))
        png = figure_bytes(render_figure(p, "light", size=(7, 5) if p["mode"] == "xyz" else (9, 3.6), dpi=170), "png")
        doc.add_picture(io.BytesIO(png), width=Inches(6.3))
        if p["caption"]:
            rtlize(doc.add_paragraph(p["caption"]))
        rows = [s for s in p["signals"] if s.get("stats")]
        if rows and p["mode"] != "xy":
            cols = [S["signal"], S["min"], S["max"], S["mean"], S["rms"], S["std"]]
            with_limits = any(s.get("limits") for s in rows)
            if with_limits:
                cols = cols + [S["limits"]]
            t = doc.add_table(rows=1, cols=len(cols))
            t.style = "Light Grid Accent 1"
            for c, name in enumerate(cols):
                t.rows[0].cells[c].text = name
            for s in rows:
                cells = t.add_row().cells
                cells[0].text = s["label"]
                for c, k in enumerate(("min", "max", "mean", "rms", "std"), start=1):
                    v = s["stats"].get(k)
                    cells[c].text = "" if v is None else f"{v:.6g}"
                if with_limits:
                    cells[len(cols) - 1].text = _limit_text(s, S)
            for row in t.rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        for r in par.runs:
                            r.font.size = Pt(8.5)
    if req.get("summary"):
        rtlize(doc.add_heading(S["summary"], 1))
        for para in req["summary"].split("\n"):
            if para.strip():
                rtlize(doc.add_paragraph(para))
    notes = sorted(req.get("notes") or [], key=lambda n: n.get("t") or 0)
    if notes:
        rtlize(doc.add_heading(S["notes"], 1))
        for nt in notes:
            rtlize(doc.add_paragraph(f"{_num(nt.get('t'))} s  {nt.get('text', '')}", style="List Bullet"))
    doc.save(path)


def build_image(store, spec: Dict, fmt: str, theme: str) -> bytes:
    p = plot_data(store, spec, for_html=False, points=4000)
    fig = render_figure(p, theme, size=(8, 6) if p.get("mode") == "xyz" else (10, 4), dpi=200 if fmt == "png" else 100)
    if p["title"]:
        fig.suptitle(_bidi(p["title"]), fontsize=11, weight="bold", color=FIG_THEME[theme]["fg"])
        fig.subplots_adjust(top=0.88)
    return figure_bytes(fig, fmt)
