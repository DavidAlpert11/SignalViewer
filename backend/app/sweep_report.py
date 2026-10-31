"""A Sweeps study as a report - the question, how the answer was reached, the best parameter sets
and the figures - in four formats:

  html  one self-contained page, interactive: the parallel plot (drag on an axis to keep a range,
        click an axis's name to colour by it), the table follows; works offline, no app needed
  docx  Word, with the figures as images and the tables as tables
  pdf   the same, page by page
  csv   the answer table itself: every parameter set with everything computed for it

The page sends what only it knows, already in the user's language: the sections (the question and
the process - filters, keep / exclude steps, merged tables, runs left out), the labels of the
columns, the axes on screen and the sets shown. The numbers and the figures are made here from the
study's answer. Figures need matplotlib (the "reports" extra); without it HTML and CSV still work.
"""
import base64
import csv
import io
import json
from datetime import datetime
from typing import Dict, List, Optional

import numpy as np

from app.report import ReportNeedsExtra, _bidi, _pyplot, missing
from svcore import sweep

HOLDS, FAILS, FAINT = "#e8961e", "#8a96a6", "#d5dbe3"


# ----------------------------------------------------------------------------- the content
def _fmt(v) -> str:
    if v is None:
        return "–"
    if isinstance(v, (float, np.floating)):
        if not np.isfinite(v):
            return "–"
        a = abs(float(v))
        return f"{v:.3g}" if (a and (a >= 1e6 or a < 1e-3)) else f"{float(v):.6g}"
    return str(v)


def _order(a: Dict) -> List[int]:
    """The sets best first: ranked ones by rank, then holding, then the rest."""
    d = a["data"]
    rank, holds = d["#rank"], d["#holds"]
    idx = list(range(len(holds)))
    return sorted(idx, key=lambda i: (rank[i] if np.isfinite(rank[i]) else np.inf, -holds[i], i))


def _value(a: Dict, name: str, i: int):
    v = a["data"][name][i]
    return v.item() if hasattr(v, "item") else v


def table(a: Dict, columns: List[str], labels: Dict[str, str], top: Optional[int]):
    cols = [c for c in columns if c in a["data"]]
    rows = []
    for i in _order(a)[: top or None]:
        rows.append([_fmt(_value(a, c, i)) if c != "#holds" else ("✓" if a["data"][c][i] == 1 else "✗") for c in cols])
    return [labels.get(c, c) for c in cols], rows


# ----------------------------------------------------------------------------- figures
def _norm(values) -> np.ndarray:
    v = np.asarray(values)
    if v.dtype.kind in "fiu":
        v = v.astype(float)
        lo, hi = np.nanmin(v) if np.isfinite(v).any() else 0, np.nanmax(v) if np.isfinite(v).any() else 1
        return (v - lo) / (hi - lo) if hi > lo else np.full(len(v), 0.5)
    labels = sorted({str(x) for x in v})
    at = {x: i for i, x in enumerate(labels)}
    return np.array([(at[str(x)] + 0.5) / len(labels) for x in v])


def fig_parallel(a: Dict, axes: List[str], labels: Dict[str, str], shown: Optional[List[int]]):
    plt = _pyplot()
    axes = [x for x in axes if x in a["data"]][:12]
    fig, ax = plt.subplots(figsize=(10, 4.2), dpi=150)
    if len(axes) < 2:
        ax.axis("off"); return fig
    ys = np.vstack([_norm(a["data"][x]) for x in axes])
    xs = np.arange(len(axes))
    holds = a["data"]["#holds"] == 1
    show = np.zeros(ys.shape[1], bool)
    show[list(shown) if shown is not None else slice(None)] = True
    alpha = 0.25 if ys.shape[1] > 2000 else 0.6
    for i in np.where(~show)[0]:
        ax.plot(xs, ys[:, i], color=FAINT, lw=0.6, alpha=0.5)
    for i in np.where(show & ~holds)[0]:
        ax.plot(xs, ys[:, i], color=FAILS, lw=0.8, alpha=alpha)
    for i in np.where(show & holds)[0]:
        ax.plot(xs, ys[:, i], color=HOLDS, lw=1.0, alpha=alpha + 0.2)
    for x in xs:
        ax.axvline(x, color="#556", lw=0.8)
    for k, name in enumerate(axes):
        v = np.asarray(a["data"][name])
        if v.dtype.kind in "fiu" and np.isfinite(v.astype(float)).any():
            ax.text(k, -0.04, _fmt(np.nanmin(v.astype(float))), ha="center", va="top", fontsize=6.5, color="#556")
            ax.text(k, 1.04, _fmt(np.nanmax(v.astype(float))), ha="center", va="bottom", fontsize=6.5, color="#556")
    ax.set_xticks(xs)
    ax.set_xticklabels([_bidi(labels.get(x, x)) for x in axes], fontsize=7, rotation=20, ha="right")
    ax.set_yticks([]); ax.set_ylim(-0.1, 1.1)
    for s in ax.spines.values():
        s.set_visible(False)
    fig.tight_layout()
    return fig


def fig_effects(eff: Dict, goal_label: str):
    plt = _pyplot()
    ps = sorted(eff["params"], key=lambda x: -x["importance"])[:6]
    n = max(1, len(ps))
    cols = min(3, n); rows = (n + cols - 1) // cols
    fig, axs = plt.subplots(rows, cols, figsize=(3.4 * cols, 2.4 * rows), dpi=150, squeeze=False)
    for k, ax in enumerate(axs.flat):
        if k >= n:
            ax.axis("off"); continue
        p = ps[k]
        xs = np.arange(len(p["levels"]))
        mean = [l["mean"] if l["mean"] is not None else np.nan for l in p["levels"]]
        lo = [l["min"] if l["min"] is not None else np.nan for l in p["levels"]]
        hi = [l["max"] if l["max"] is not None else np.nan for l in p["levels"]]
        ax.fill_between(xs, lo, hi, color="#9cc0f5", alpha=0.35, lw=0)
        ax.plot(xs, mean, color="#2f6fd0", marker="o", ms=3, lw=1.5)
        ax.set_xticks(xs); ax.set_xticklabels([_fmt(l["value"]) for l in p["levels"]], fontsize=6.5)
        ax.set_title(_bidi(f"{p['param']} · {round(p['importance'] * 100)}%"), fontsize=8)
        ax.tick_params(labelsize=6.5)
    fig.suptitle(_bidi(goal_label), fontsize=8)
    fig.tight_layout()
    return fig


def fig_map(m: Dict):
    plt = _pyplot()
    share = np.array([[np.nan if v is None else v for v in r] for r in m["share"]], dtype=float)
    fig, ax = plt.subplots(figsize=(1.2 + 0.9 * len(m["xs"]), 1.0 + 0.6 * len(m["ys"])), dpi=150)
    ax.imshow(share, origin="lower", cmap="viridis", vmin=0, vmax=1, aspect="auto")
    for j, r in enumerate(m["best"]):
        for i, b in enumerate(r):
            if b is not None:
                ax.text(i, j, _fmt(b), ha="center", va="center", fontsize=6, color="white")
    ax.set_xticks(range(len(m["xs"]))); ax.set_xticklabels([_fmt(x) for x in m["xs"]], fontsize=7)
    ax.set_yticks(range(len(m["ys"]))); ax.set_yticklabels([_fmt(y) for y in m["ys"]], fontsize=7)
    ax.set_xlabel(_bidi(m["x"]), fontsize=8); ax.set_ylabel(_bidi(m["y"]), fontsize=8)
    fig.tight_layout()
    return fig


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    _pyplot().close(fig)
    return buf.getvalue()


def figures(p, q: Dict, a: Dict, req: Dict) -> Dict[str, bytes]:
    """The figures asked for, as PNG; none (and no error) without matplotlib."""
    inc = req.get("include") or {}
    out: Dict[str, bytes] = {}
    labels = req.get("labels") or {}
    try:
        if inc.get("plot", True):
            out["plot"] = _png(fig_parallel(a, req.get("axes") or [], labels, req.get("shown")))
        if inc.get("effects", True):
            val = a["data"]["#goal"] if "#goal" in a["data"] else a["data"]["#holds"]
            out["effects"] = _png(fig_effects(sweep.effects(p, val, a["data"]["#holds"]), labels.get("#effects", "")))
        mx, my = req.get("mapX"), req.get("mapY")
        if inc.get("map", True) and mx in p.params and my in p.params and mx != my:
            val = a["data"]["#goal"] if "#goal" in a["data"] else a["data"]["#holds"]
            out["map"] = _png(fig_map(sweep.pair_map(p, mx, my, val, a["data"]["#holds"])))
    except ReportNeedsExtra:
        if req.get("format") in ("docx", "pdf"):
            raise
    return out


# ----------------------------------------------------------------------------- formats
def build(p, q: Dict, a: Dict, req: Dict) -> str:
    fmt, path = req["format"], req["path"]
    if fmt == "csv":
        return _csv(a, req, path)
    figs = figures(p, q, a, req)
    if fmt == "html":
        with open(path, "w", encoding="utf-8") as f:
            f.write(_html(a, req, figs))
    elif fmt == "docx":
        _docx(a, req, figs, path)
    elif fmt == "pdf":
        _pdf(a, req, figs, path)
    else:
        raise ValueError(f"Unknown format: {fmt}")
    return path


def _csv(a: Dict, req: Dict, path: str) -> str:
    labels = req.get("labels") or {}
    names = [m["name"] for m in a["columns"] if m["name"] not in ("#complete",)]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([labels.get(n, n) for n in names])
        for i in _order(a):
            w.writerow(["" if (isinstance(v := _value(a, n, i), float) and not np.isfinite(v)) else v for n in names])
    return path


def _sections(req: Dict) -> List[Dict]:
    return [s for s in (req.get("sections") or []) if s.get("items") or s.get("text")]


def _docx(a: Dict, req: Dict, figs: Dict[str, bytes], path: str):
    try:
        import docx  # noqa: F401
    except ImportError:
        raise ReportNeedsExtra(missing("Word reports", "reports")) from None
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    rtl = req.get("lang") == "he"
    L = req.get("labels") or {}
    doc = Document()

    def r(par):
        if rtl:
            ppr = par._p.get_or_add_pPr()
            b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); ppr.append(b)
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in par.runs:
                rt = OxmlElement("w:rtl"); rt.set(qn("w:val"), "1"); run._r.get_or_add_rPr().append(rt)
        return par

    def tbl(head, rows):
        t = doc.add_table(rows=1, cols=max(1, len(head)))
        t.style = "Light Grid Accent 1"
        if rtl:
            tp = t._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); bv.set(qn("w:val"), "1"); tp.append(bv)
        for c, h in enumerate(head):
            t.rows[0].cells[c].text = str(h)
        for row in rows:
            cells = t.add_row().cells
            for c, v in enumerate(row):
                cells[c].text = str(v)

    r(doc.add_heading(req.get("title") or "Sweeps", 0))
    r(doc.add_paragraph(f"{req.get('subtitle', '')} · {datetime.now().strftime('%Y-%m-%d %H:%M')}"))
    for para in (req.get("notes") or "").split("\n"):
        if para.strip():
            r(doc.add_paragraph(para))
    r(doc.add_heading(L.get("#answer", "Answer"), 1))
    tbl([L.get("#what", ""), L.get("#value", "")], req.get("summary") or [])
    for sec in _sections(req):
        r(doc.add_heading(str(sec.get("title") or ""), 1))
        for it in sec.get("items") or []:
            r(doc.add_paragraph(str(it), style="List Bullet"))
    for key, title in (("plot", L.get("#plot", "")), ("effects", L.get("#effectsTitle", "")), ("map", L.get("#mapTitle", ""))):
        if key in figs:
            r(doc.add_heading(title, 1))
            doc.add_picture(io.BytesIO(figs[key]), width=Inches(6.5))
    if (req.get("include") or {}).get("table", True):
        head, rows = table(a, req.get("columns") or [], L, req.get("top") or 25)
        r(doc.add_heading(L.get("#best", "Best"), 1))
        tbl(head, rows)
    doc.save(path)


def _pdf(a: Dict, req: Dict, figs: Dict[str, bytes], path: str):
    plt = _pyplot()
    from matplotlib.backends.backend_pdf import PdfPages
    import matplotlib.image as mpimg
    rtl = req.get("lang") == "he"
    L = req.get("labels") or {}
    X, ha = (0.92, "right") if rtl else (0.08, "left")

    def text_page(lines: List[tuple]):
        fig = plt.figure(figsize=(8.27, 11.69))
        y = 0.94
        for size, s in lines:
            if y < 0.06:
                pdf.savefig(fig); plt.close(fig); fig = plt.figure(figsize=(8.27, 11.69)); y = 0.94
            fig.text(X, y, _bidi(s), ha=ha, va="top", fontsize=size, wrap=True)
            y -= 0.012 * size / 6 + 0.012
        pdf.savefig(fig); plt.close(fig)

    with PdfPages(path) as pdf:
        lines = [(18, req.get("title") or "Sweeps"), (9, f"{req.get('subtitle', '')} · {datetime.now().strftime('%Y-%m-%d %H:%M')}")]
        lines += [(10, x) for x in (req.get("notes") or "").split("\n") if x.strip()]
        lines += [(13, L.get("#answer", "Answer"))] + [(10, f"{k}: {v}") for k, v in (req.get("summary") or [])]
        for sec in _sections(req):
            lines += [(13, str(sec.get("title") or ""))] + [(9.5, f"•  {it}") for it in sec.get("items") or []]
        text_page(lines)
        for key, title in (("plot", L.get("#plot", "")), ("effects", L.get("#effectsTitle", "")), ("map", L.get("#mapTitle", ""))):
            if key in figs:
                fig = plt.figure(figsize=(11.69, 8.27))
                fig.text(X, 0.95, _bidi(title), ha=ha, fontsize=14)
                ax = fig.add_axes([0.04, 0.05, 0.92, 0.85]); ax.axis("off")
                ax.imshow(mpimg.imread(io.BytesIO(figs[key]), format="png"))
                pdf.savefig(fig); plt.close(fig)
        if (req.get("include") or {}).get("table", True):
            head, rows = table(a, req.get("columns") or [], L, req.get("top") or 25)
            for start in range(0, max(1, len(rows)), 30):
                fig = plt.figure(figsize=(11.69, 8.27))
                fig.text(X, 0.95, _bidi(L.get("#best", "Best")), ha=ha, fontsize=14)
                ax = fig.add_axes([0.03, 0.04, 0.94, 0.86]); ax.axis("off")
                part = rows[start:start + 30] or [["–"] * len(head)]
                t = ax.table(cellText=part, colLabels=[_bidi(h) for h in head], loc="upper center", cellLoc="center")
                t.auto_set_font_size(False); t.set_fontsize(7); t.scale(1, 1.25)
                pdf.savefig(fig); plt.close(fig)


# ----------------------------------------------------------------------------- interactive HTML
def _col_json(a: Dict, name: str) -> Dict:
    v = np.asarray(a["data"][name])
    if v.dtype.kind in "fiu":
        f = v.astype(float)
        return {"kind": "num", "v": [None if not np.isfinite(x) else round(float(x), 6) for x in f]}
    labels = sorted({str(x) for x in v})
    try:
        labels = sorted(labels, key=float)
    except ValueError:
        pass
    at = {x: i for i, x in enumerate(labels)}
    return {"kind": "cat", "labels": labels, "v": [at[str(x)] for x in v]}


def _esc(s: str) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _html(a: Dict, req: Dict, figs: Dict[str, bytes]) -> str:
    L = req.get("labels") or {}
    lang = req.get("lang", "en")
    names = [m["name"] for m in a["columns"] if m["name"] not in ("#complete", "#partial")]
    data = {"cols": {n: {**_col_json(a, n), "label": L.get(n, n)} for n in names},
            "axes": [x for x in (req.get("axes") or []) if x in a["data"]],
            "table": [c for c in (req.get("columns") or []) if c in a["data"]],
            "n": int(len(a["data"]["#holds"])), "ui": {k: L.get(k, k) for k in ("#shown", "#clear", "#hint", "#best", "#colour", "#holdsKey")}}
    sec = "".join(f"<h2>{_esc(s.get('title', ''))}</h2><ul>{''.join(f'<li>{_esc(x)}</li>' for x in s.get('items') or [])}</ul>"
                  for s in _sections(req))
    summary = "".join(f"<tr><td>{_esc(k)}</td><td>{_esc(v)}</td></tr>" for k, v in (req.get("summary") or []))
    imgs = "".join(f"<h2>{_esc(L.get(t, ''))}</h2><img alt='' src='data:image/png;base64,{base64.b64encode(figs[k]).decode()}'>"
                   for k, t in (("effects", "#effectsTitle"), ("map", "#mapTitle")) if k in figs)
    notes = "".join(f"<p>{_esc(x)}</p>" for x in (req.get("notes") or "").split("\n") if x.strip())
    return (_PAGE.replace("{{lang}}", lang).replace("{{dir}}", "rtl" if lang == "he" else "ltr")
            .replace("{{title}}", _esc(req.get("title") or "Sweeps"))
            .replace("{{subtitle}}", _esc(f"{req.get('subtitle', '')} · {datetime.now().strftime('%Y-%m-%d %H:%M')}"))
            .replace("{{notes}}", notes).replace("{{answer}}", _esc(L.get("#answer", "Answer"))).replace("{{summary}}", summary)
            .replace("{{plotTitle}}", _esc(L.get("#plot", ""))).replace("{{sections}}", sec).replace("{{images}}", imgs)
            .replace("{{best}}", _esc(L.get("#best", "")))
            .replace("{{data}}", json.dumps(data, ensure_ascii=False).replace("</", "<\\/")))


_PAGE = r"""<!doctype html><html lang="{{lang}}" dir="{{dir}}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{title}}</title><style>
body{font-family:system-ui,"Segoe UI",sans-serif;margin:28px auto;max-width:1200px;padding:0 20px;color:#1f2328}
h1{margin:0 0 4px}h2{margin:28px 0 8px;font-size:1.15rem}.muted{color:#667}
table{border-collapse:collapse;font-size:13px}th,td{border-bottom:1px solid #e3e6ea;padding:5px 9px;text-align:start;white-space:nowrap}
th{background:#f4f6f8;position:sticky;top:0}td.n{text-align:end;font-variant-numeric:tabular-nums}
.ok{color:#1a7f37}.bad{color:#c62828}img{max-width:100%;border:1px solid #e3e6ea;border-radius:8px}
#wrap{position:relative;border:1px solid #e3e6ea;border-radius:8px}canvas{display:block;width:100%;cursor:crosshair}
.ax{position:absolute;top:6px;transform:translateX(-50%);font-size:12px;border:1px solid transparent;border-radius:5px;background:#fff;
padding:1px 6px;cursor:pointer;max-width:150px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;unicode-bidi:plaintext}
.ax:hover{border-color:#ccd}.ax.on{color:#e8961e}.bar{display:flex;gap:12px;align-items:center;margin:8px 0;flex-wrap:wrap}
button.b{border:1px solid #ccd;background:#fff;border-radius:6px;padding:4px 10px;cursor:pointer}
#tw{max-height:460px;overflow:auto;border:1px solid #e3e6ea;border-radius:8px}
</style></head><body>
<h1>{{title}}</h1><p class="muted">{{subtitle}}</p>{{notes}}
<h2>{{answer}}</h2><table>{{summary}}</table>
{{sections}}
<h2>{{plotTitle}}</h2>
<div class="bar"><span id="count" class="muted"></span><button class="b" id="clear"></button><span class="muted" id="hint"></span><span id="key"></span></div>
<div id="wrap"><canvas id="cv" height="420"></canvas></div>
<h2>{{best}}</h2><div id="tw"><table id="tb"></table></div>
{{images}}
<script>
const D={{data}};
const cv=document.getElementById('cv'),wrap=document.getElementById('wrap'),ctx=cv.getContext('2d');
const T=40,B=34,PX=60,H=420,N=D.n;let W=900,filters={},colorBy='#holds',brush=null;
const rtl=document.dir==='rtl';
document.getElementById('clear').textContent=D.ui['#clear'];document.getElementById('hint').textContent=D.ui['#hint'];
function col(n){return D.cols[n]}
function scale(n){const c=col(n);if(c.kind==='cat'){const k=c.labels.length;return v=>v==null?NaN:(v+.5)/k}
 let lo=Infinity,hi=-Infinity;for(const v of c.v)if(v!=null){lo=Math.min(lo,v);hi=Math.max(hi,v)}if(!(hi>lo)){lo-=1;hi+=1}return v=>v==null?NaN:(v-lo)/(hi-lo)}
function inv(n,t){const c=col(n);let lo=Infinity,hi=-Infinity;for(const v of c.v)if(v!=null){lo=Math.min(lo,v);hi=Math.max(hi,v)}return lo+t*(hi-lo)}
function keeps(i){for(const n in filters){const v=col(n).v[i];if(v==null)return false;if(!filters[n].some(([a,b])=>v>=a&&v<=b))return false}return true}
function ax(i){const k=D.axes.length;return k<2?W/2:PX+i*(W-2*PX)/(k-1)}
function ramp(t){const S=[[68,1,84],[59,82,139],[33,145,140],[94,201,98],[253,231,37]];if(!isFinite(t))return'#999';const x=Math.max(0,Math.min(1,t))*4,i=Math.min(3,Math.floor(x)),f=x-i;return'rgb('+S[i].map((a,k)=>Math.round(a+(S[i+1][k]-a)*f)).join(',')+')'}
const PAL=['#2f6fd0','#e8961e','#8e5cd9','#1a9e6a','#d64550','#2bb0b8','#b08a2e'];
function colour(i){if(!colorBy)return'#e8961e';const c=col(colorBy),v=c.v[i];if(colorBy==='#holds'||colorBy==='#pareto')return v===1?'#e8961e':'#8a96a6';
 if(c.kind==='cat')return v==null?'#999':PAL[v%PAL.length];return ramp(scale(colorBy)(v))}
function draw(){W=wrap.clientWidth;const dpr=window.devicePixelRatio||1;cv.width=W*dpr;cv.height=H*dpr;cv.style.height=H+'px';ctx.setTransform(dpr,0,0,dpr,0,0);ctx.clearRect(0,0,W,H);
 const h=H-T-B,sc=D.axes.map(scale);const y=(k,i)=>{const t=sc[k](col(D.axes[k]).v[i]);return isFinite(t)?T+(1-t)*h:H-10};
 let shown=0;const keep=new Uint8Array(N);for(let i=0;i<N;i++)if(keeps(i)){keep[i]=1;shown++}
 ctx.globalAlpha=.15;ctx.strokeStyle='#ccd';ctx.beginPath();for(let i=0;i<N;i++)if(!keep[i]){ctx.moveTo(ax(0),y(0,i));for(let k=1;k<D.axes.length;k++)ctx.lineTo(ax(k),y(k,i))}ctx.stroke();
 ctx.globalAlpha=shown>3000?.3:shown>300?.55:.85;for(let i=0;i<N;i++)if(keep[i]){ctx.strokeStyle=colour(i);ctx.beginPath();ctx.moveTo(ax(0),y(0,i));for(let k=1;k<D.axes.length;k++)ctx.lineTo(ax(k),y(k,i));ctx.stroke()}
 ctx.globalAlpha=1;ctx.font='11px system-ui';ctx.fillStyle='#667';
 D.axes.forEach((n,k)=>{const x=ax(k);ctx.strokeStyle='#889';ctx.beginPath();ctx.moveTo(x,T);ctx.lineTo(x,T+h);ctx.stroke();const c=col(n);ctx.textAlign='right';
  if(c.kind==='cat')c.labels.forEach((l,j)=>ctx.fillText(l,x-4,T+(1-(j+.5)/c.labels.length)*h+4));else{ctx.fillText(fmt(inv(n,1)),x-4,T+8);ctx.fillText(fmt(inv(n,0)),x-4,T+h)}
  (filters[n]||[]).forEach(([a,b])=>{const s=sc[k];const y1=T+(1-s(a))*h,y2=T+(1-s(b))*h;ctx.fillStyle='rgba(232,150,30,.25)';ctx.fillRect(x-6,Math.min(y1,y2),12,Math.abs(y2-y1)||2);ctx.fillStyle='#667'});
  if(brush&&brush.k===k){ctx.fillStyle='rgba(47,111,208,.25)';ctx.fillRect(x-7,Math.min(brush.y0,brush.y1),14,Math.abs(brush.y1-brush.y0));ctx.fillStyle='#667'}});
 document.getElementById('count').textContent=D.ui['#shown'].replace('{n}',shown).replace('{of}',N);
 names();rows(keep)}
function fmt(v){if(v==null||!isFinite(v))return'–';const a=Math.abs(v);return a&&(a>=1e6||a<1e-3)?v.toExponential(2):String(+v.toPrecision(4))}
function names(){wrap.querySelectorAll('.ax').forEach(e=>e.remove());D.axes.forEach((n,k)=>{const b=document.createElement('button');b.className='ax'+(colorBy===n?' on':'');b.textContent=col(n).label;b.title=D.ui['#colour'];
 b.style.left=ax(k)+'px';b.onclick=()=>{colorBy=colorBy===n?'#holds':n;draw()};wrap.appendChild(b)})}
function cell(n,i){const c=col(n),v=c.v[i];if(n==='#holds')return v===1?'<td class="ok">✓</td>':'<td class="bad">✗</td>';if(c.kind==='cat')return'<td>'+(v==null?'–':c.labels[v])+'</td>';return'<td class="n">'+fmt(v)+'</td>'}
function rows(keep){const r=D.cols['#rank'].v;const idx=[];for(let i=0;i<N;i++)if(keep[i])idx.push(i);
 idx.sort((a,b)=>(r[a]??1e18)-(r[b]??1e18)||D.cols['#holds'].v[b]-D.cols['#holds'].v[a]);
 const cols=D.table.includes('#holds')?D.table:[...D.table,'#holds'];document.getElementById('tb').innerHTML='<tr>'+cols.map(n=>'<th>'+col(n).label+'</th>').join('')+'</tr>'+idx.slice(0,300).map(i=>'<tr>'+cols.map(n=>cell(n,i)).join('')+'</tr>').join('')}
cv.onpointerdown=e=>{const r=cv.getBoundingClientRect(),x=e.clientX-r.left,y=e.clientY-r.top;for(let k=0;k<D.axes.length;k++)if(Math.abs(x-ax(k))<12){brush={k,y0:y,y1:y,add:e.shiftKey};cv.setPointerCapture(e.pointerId);return}};
cv.onpointermove=e=>{if(!brush)return;const r=cv.getBoundingClientRect();brush.y1=Math.max(T,Math.min(H-B,e.clientY-r.top));draw()};
cv.onpointerup=()=>{if(!brush)return;const b=brush;brush=null;const n=D.axes[b.k],h=H-T-B;if(Math.abs(b.y1-b.y0)<3){delete filters[n];draw();return}
 const t0=1-(Math.min(b.y0,b.y1)-T)/h,t1=1-(Math.max(b.y0,b.y1)-T)/h,c=col(n);let rg;
 if(c.kind==='cat'){const k=c.labels.length;const lo=Math.ceil(t1*k-.5),hi=Math.floor(t0*k-.5);rg=[lo,hi]}else rg=[inv(n,t1),inv(n,t0)];
 filters[n]=b.add&&filters[n]?[...filters[n],rg]:[rg];draw()};
document.getElementById('clear').onclick=()=>{filters={};draw()};window.onresize=draw;draw();
</script></body></html>"""
